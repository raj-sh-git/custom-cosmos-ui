import os
import io
import csv
import json
import uuid
import time
import queue
import random
import threading
import traceback
from datetime import datetime, timezone, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed
from flask import (
    Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_file, Blueprint, Response, abort, has_request_context
)
from flask_session import Session
from werkzeug.security import generate_password_hash, check_password_hash
from azure.identity import ClientSecretCredential
from azure.cosmos import CosmosClient, PartitionKey, exceptions as cosmos_exceptions
from openpyxl import Workbook, load_workbook
from werkzeug.middleware.proxy_fix import ProxyFix

app = Flask(__name__, static_url_path="/cosmos-ui/static")
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)

# Server-side Session Configuration
SESSION_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "flask_sessions")
os.makedirs(SESSION_DIR, exist_ok=True)

TEMP_UPLOAD_DIR = os.path.join(SESSION_DIR, "temp_imports")
os.makedirs(TEMP_UPLOAD_DIR, exist_ok=True)

TEMP_EXPORT_DIR = os.path.join(SESSION_DIR, "temp_exports")
os.makedirs(TEMP_EXPORT_DIR, exist_ok=True)

app.secret_key = os.environ.get("FLASK_SECRET_KEY", str(uuid.uuid4()))
app.config.update(
    SESSION_TYPE="filesystem",
    SESSION_FILE_DIR=SESSION_DIR,
    SESSION_PERMANENT=True,
    PERMANENT_SESSION_LIFETIME=86400, # 24 hours
)
Session(app)

# System Database and Container Configuration (Overridable via Environment Variables)
SYSTEM_DB = os.environ.get("COSMOS_AUTH_DB", os.environ.get("SYSTEM_DB", "cosmos-access"))
USER_CONTAINER = os.environ.get("COSMOS_USER_CONTAINER", os.environ.get("USER_CONTAINER", "cosmosusers"))
LOGS_CONTAINER = os.environ.get("COSMOS_LOGS_CONTAINER", os.environ.get("LOGS_CONTAINER", "cosmosactivitylogs"))
USER_PK_PATH = os.environ.get("COSMOS_USER_PK_PATH", "/id")
LOGS_PK_PATH = os.environ.get("COSMOS_LOGS_PK_PATH", "/id")
DEFAULT_PASSWORD_EXPIRY_DAYS = int(os.environ.get("DEFAULT_PASSWORD_EXPIRY_DAYS", "90"))
ACTIVITY_LOG_RETENTION_DAYS = int(os.environ.get("ACTIVITY_LOG_RETENTION_DAYS", "30"))

@app.context_processor
def inject_user_context():
    user = session.get('user') if (has_request_context() and 'user' in session) else None
    return {
        'current_user': user,
        'user_role': (user.get('role') if user else 'contributor'),
        'is_admin': (user.get('role') == 'admin') if user else False,
        'is_reader': (user.get('role') == 'reader') if user else False,
        'system_db_name': SYSTEM_DB,
        'user_container_name': USER_CONTAINER,
        'logs_container_name': LOGS_CONTAINER
    }

# In-memory store for active CosmosClient instances
CLIENT_STORE = {}

# Background Ingestion Tasks Store
IMPORT_TASKS = {}
IMPORT_TASKS_LOCK = threading.Lock()

# Background Export Tasks Store
EXPORT_TASKS = {}
EXPORT_TASKS_LOCK = threading.Lock()

# ---------- Blueprint ----------
ui = Blueprint("ui", __name__, url_prefix="/cosmos-ui")

def get_store():
    sid = session.get("sid")
    return CLIENT_STORE.get(sid) if sid else None

def auto_connect_from_env_if_available():
    """Auto-connects to Cosmos DB if environment variables are set."""
    if get_store():
        return True

    conn_str = os.environ.get("AZURE_COSMOS_CONNECTION_STRING") or os.environ.get("COSMOS_CONNECTION_STRING")
    if conn_str:
        try:
            client = CosmosClient.from_connection_string(conn_str)
            _ = list(client.list_databases())
            sid = session.get("sid") or str(uuid.uuid4())
            session["sid"] = sid
            endpoint_name = "Cosmos DB Account"
            for part in conn_str.split(";"):
                if part.startswith("AccountEndpoint="):
                    endpoint_name = part.replace("AccountEndpoint=", "")
            session["auth_info"] = {
                "endpoint": endpoint_name,
                "method": "Connection String (Env)"
            }
            CLIENT_STORE[sid] = {
                "client": client,
                "login_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            return True
        except Exception as e:
            print(f"Env conn_str connection notice: {e}")

    endpoint = os.environ.get("AZURE_COSMOS_ENDPOINT") or os.environ.get("COSMOS_ENDPOINT")
    key = os.environ.get("AZURE_COSMOS_KEY") or os.environ.get("COSMOS_KEY")
    if endpoint and key:
        try:
            client = CosmosClient(endpoint, credential=key, connection_verify=True)
            _ = list(client.list_databases())
            sid = session.get("sid") or str(uuid.uuid4())
            session["sid"] = sid
            session["auth_info"] = {
                "endpoint": endpoint,
                "method": "Account Key (Env)"
            }
            CLIENT_STORE[sid] = {
                "client": client,
                "login_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            return True
        except Exception as e:
            print(f"Env endpoint/key connection notice: {e}")

    return False

def get_cosmos_client():
    store = get_store()
    if store:
        return store["client"]
    if auto_connect_from_env_if_available():
        store = get_store()
        return store["client"] if store else None
    return None

def is_cosmos_connected():
    return bool(get_cosmos_client())

def require_auth():
    """Requires both Cosmos DB backend connection and valid user login session."""
    if not is_cosmos_connected():
        return False
    user = session.get("user")
    return bool(user and user.get("is_authenticated"))

def is_admin():
    return require_auth() and session.get("user", {}).get("role") == "admin"

def has_write_permission():
    return require_auth() and session.get("user", {}).get("role") in ["admin", "contributor"]

def login_required(f):
    from functools import wraps
    @wraps(f)
    def inner(*a, **kw):
        if not require_auth():
            flash("Please sign in to access this resource.", "warning")
            return redirect(url_for("ui.login"))
        return f(*a, **kw)
    return inner

def admin_required(f):
    from functools import wraps
    @wraps(f)
    def inner(*a, **kw):
        if not require_auth():
            flash("Please sign in to access this resource.", "warning")
            return redirect(url_for("ui.login"))
        if not is_admin():
            flash("Permission denied: Administrator role required.", "danger")
            return redirect(url_for("ui.dashboard"))
        return f(*a, **kw)
    return inner

def write_permission_required(f):
    from functools import wraps
    @wraps(f)
    def inner(*a, **kw):
        if not require_auth():
            flash("Please sign in to access this resource.", "warning")
            return redirect(url_for("ui.login"))
        if not has_write_permission():
            flash("Permission denied: Read-only accounts cannot modify resources.", "danger")
            return redirect(request.referrer or url_for("ui.dashboard"))
        return f(*a, **kw)
    return inner

# ---------- System Database & User Management Helpers ----------

def normalize_pk_path(pk):
    """Ensures partition key path starts with a leading slash."""
    if not pk:
        return "/id"
    pk = str(pk).strip()
    return pk if pk.startswith("/") else "/" + pk

def ensure_system_database(client):
    """Idempotently ensures SYSTEM_DB exists in Cosmos DB."""
    try:
        db = client.get_database_client(SYSTEM_DB)
        db.read()
        return db
    except Exception:
        pass

    try:
        return client.create_database_if_not_exists(id=SYSTEM_DB)
    except Exception as e1:
        try:
            return client.create_database(id=SYSTEM_DB)
        except Exception as e2:
            print(f"Notice: database creation fallback for {SYSTEM_DB}: {e2}")
            return client.get_database_client(SYSTEM_DB)

def ensure_system_container(client, container_name, pk_path):
    """
    Idempotently creates the system container with multi-tier fallback for serverless,
    provisioned throughput (400 RU/s), and shared database throughput modes.
    """
    db = ensure_system_database(client)
    norm_pk = normalize_pk_path(pk_path)
    pk = PartitionKey(path=norm_pk)

    # 1. Check if container already exists
    try:
        c = db.get_container_client(container_name)
        c.read()
        return c
    except Exception:
        pass

    # 2. Try creating container without explicit throughput (serverless or shared-throughput DB)
    try:
        return db.create_container_if_not_exists(id=container_name, partition_key=pk)
    except Exception as e_no_tp:
        print(f"Notice: create_container_if_not_exists ({container_name}) without throughput: {e_no_tp}")

    # 3. Try creating container with default 400 RU/s (provisioned non-shared DB)
    try:
        return db.create_container_if_not_exists(id=container_name, partition_key=pk, offer_throughput=400)
    except Exception as e_tp:
        print(f"Notice: create_container_if_not_exists ({container_name}) with 400 RU/s: {e_tp}")

    # 4. Try direct create_container without throughput
    try:
        return db.create_container(id=container_name, partition_key=pk)
    except Exception as e_dir:
        print(f"Notice: create_container ({container_name}) direct: {e_dir}")

    # 5. Try direct create_container with 400 RU/s
    try:
        return db.create_container(id=container_name, partition_key=pk, offer_throughput=400)
    except Exception as e_final:
        print(f"Error: All creation attempts for system container {container_name} failed: {e_final}")
        return db.get_container_client(container_name)

def get_system_user_container(client=None):
    """Returns the container client for user management, ensuring DB and container exist."""
    cl = client or get_cosmos_client()
    if not cl:
        raise ValueError("Cosmos DB client is not initialized.")
    return ensure_system_container(cl, USER_CONTAINER, USER_PK_PATH)

def get_system_logs_container(client=None):
    """Returns the container client for activity logs, ensuring DB and container exist."""
    cl = client or get_cosmos_client()
    if not cl:
        raise ValueError("Cosmos DB client is not initialized.")
    return ensure_system_container(cl, LOGS_CONTAINER, LOGS_PK_PATH)

def ensure_system_containers_exist(client=None):
    """Ensures SYSTEM_DB, USER_CONTAINER, and LOGS_CONTAINER exist in Cosmos DB."""
    try:
        cl = client or get_cosmos_client()
        if not cl:
            return
        get_system_user_container(cl)
        get_system_logs_container(cl)
    except Exception as e:
        print(f"Notice: system containers check/create: {e}")
        traceback.print_exc()

def check_has_any_users(client=None):
    """Checks whether any user accounts exist in USER_CONTAINER."""
    try:
        cl = client or get_cosmos_client()
        if not cl:
            return False
        container = get_system_user_container(cl)
        res = list(container.query_items("SELECT VALUE COUNT(1) FROM c", enable_cross_partition_query=True))
        return bool(res and res[0] > 0)
    except cosmos_exceptions.CosmosResourceNotFoundError:
        return False
    except Exception as e:
        print(f"Notice: error checking users in {USER_CONTAINER}: {e}")
        return False

def get_user_by_username(client, username):
    """Retrieves a user entity by username (case-insensitive)."""
    if not username or not client:
        return None
    try:
        uname = username.strip().lower()
        container = get_system_user_container(client)
        query = "SELECT * FROM c WHERE LOWER(c.username) = @uname OR c.id = @uname"
        params = [{"name": "@uname", "value": uname}]
        items = list(container.query_items(query=query, parameters=params, enable_cross_partition_query=True))
        return items[0] if items else None
    except cosmos_exceptions.CosmosResourceNotFoundError:
        return None
    except Exception as e:
        print(f"Error fetching user {username}: {e}")
        return None

def save_user(client, username, password=None, password_hash=None, email=None, display_name=None, role=None, is_active=None, must_change_password=None, password_expiry_days=None, update_login=False):
    """
    Creates or updates a user document in USER_CONTAINER with salted one-way scrypt password hashing and password expiry tracking.
    """
    uname = username.strip()
    row_id = uname.lower()
    container = get_system_user_container(client)

    existing = get_user_by_username(client, uname)
    now_iso = datetime.now(timezone.utc).isoformat()

    if existing is None:
        user_doc = {
            "id": row_id,
            "username": uname,
            "display_name": (display_name or uname).strip(),
            "email": (email or "").strip(),
            "role": (role or "contributor").lower().strip(),
            "is_active": True if is_active is None else bool(is_active),
            "must_change_password": False if must_change_password is None else bool(must_change_password),
            "password_expiry_days": int(password_expiry_days if password_expiry_days is not None else DEFAULT_PASSWORD_EXPIRY_DAYS),
            "password_last_set_at": now_iso,
            "created_at": now_iso,
            "last_login": now_iso if update_login else ""
        }
    else:
        user_doc = dict(existing)
        if email is not None:
            user_doc["email"] = email.strip()
        if display_name is not None:
            user_doc["display_name"] = display_name.strip()
        if role is not None:
            user_doc["role"] = role.lower().strip()
        if is_active is not None:
            user_doc["is_active"] = bool(is_active)
        if must_change_password is not None:
            user_doc["must_change_password"] = bool(must_change_password)
        if password_expiry_days is not None:
            try:
                user_doc["password_expiry_days"] = int(password_expiry_days)
            except (ValueError, TypeError):
                user_doc["password_expiry_days"] = DEFAULT_PASSWORD_EXPIRY_DAYS
        if update_login:
            user_doc["last_login"] = now_iso

    if password:
        user_doc["password_hash"] = generate_password_hash(password, method="scrypt")
        user_doc["password_last_set_at"] = now_iso
    elif password_hash:
        user_doc["password_hash"] = password_hash
        user_doc["password_last_set_at"] = now_iso

    try:
        execute_with_429_retry(container.upsert_item, body=user_doc)
    except cosmos_exceptions.CosmosResourceNotFoundError:
        container = ensure_system_container(client, USER_CONTAINER, USER_PK_PATH)
        execute_with_429_retry(container.upsert_item, body=user_doc)

    return user_doc

def get_all_users(client):
    """List all registered users from USER_CONTAINER with password expiry calculations."""
    try:
        container = get_system_user_container(client)
        items = list(container.query_items("SELECT * FROM c", enable_cross_partition_query=True))
        users = []
        for u in items:
            u.setdefault("display_name", u.get("username", ""))
            u.setdefault("email", "")
            u.setdefault("role", "contributor")
            u.setdefault("is_active", True)
            u.setdefault("must_change_password", False)
            u.setdefault("password_expiry_days", DEFAULT_PASSWORD_EXPIRY_DAYS)
            u.setdefault("password_last_set_at", u.get("created_at", ""))
            u.setdefault("last_login", "")
            u.setdefault("created_at", "")

            try:
                expiry_days = int(u.get("password_expiry_days", DEFAULT_PASSWORD_EXPIRY_DAYS) or 0)
            except Exception:
                expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS
            u["password_expiry_days"] = expiry_days

            last_set = u.get("password_last_set_at") or u.get("created_at")
            if expiry_days <= 0:
                u["expiry_status"] = "Never Expires"
                u["days_remaining"] = None
                u["is_expired"] = False
            elif last_set:
                try:
                    last_set_dt = datetime.fromisoformat(last_set.replace("Z", "+00:00"))
                    age_days = (datetime.now(timezone.utc) - last_set_dt).days
                    rem_days = expiry_days - age_days
                    u["days_remaining"] = rem_days
                    u["password_age_days"] = age_days
                    if rem_days <= 0:
                        u["expiry_status"] = "Expired"
                        u["is_expired"] = True
                    else:
                        u["expiry_status"] = f"{rem_days} days left"
                        u["is_expired"] = False
                except Exception:
                    u["expiry_status"] = f"{expiry_days} days"
                    u["days_remaining"] = expiry_days
                    u["is_expired"] = False
            else:
                u["expiry_status"] = f"{expiry_days} days"
                u["days_remaining"] = expiry_days
                u["is_expired"] = False

            users.append(u)
        users.sort(key=lambda x: x.get("created_at", "") or x.get("username", ""))
        return users
    except cosmos_exceptions.CosmosResourceNotFoundError:
        return []
    except Exception as e:
        print(f"Error listing users: {e}")
        return []

def delete_user_by_username(client, username):
    """Deletes a user document by username."""
    uname = username.strip().lower()
    container = get_system_user_container(client)
    user = get_user_by_username(client, uname)
    if user:
        item_id = user["id"]
        pk_path = normalize_pk_path(USER_PK_PATH)
        pk_key = pk_path.strip("/")
        pk_val = user.get(pk_key, item_id)
        execute_with_429_retry(container.delete_item, item=item_id, partition_key=pk_val)

def parse_users_file(client, file_obj, filename):
    """
    Parses a CSV or Excel (.xlsx) file, validates fields, checks existing users in Cosmos DB,
    and returns parsed list with masked passwords, summary counts, and error messages.
    """
    rows = []
    fn = filename.lower()
    errors = []

    try:
        if fn.endswith(".csv"):
            content = file_obj.read().decode("utf-8-sig", errors="ignore")
            reader = csv.DictReader(io.StringIO(content))
            for r in reader:
                rows.append(r)
        elif fn.endswith(".xlsx"):
            wb = load_workbook(file_obj, data_only=True)
            ws = wb.active
            headers = None
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx == 0:
                    headers = [str(h).strip().lower() if h is not None else "" for h in row]
                else:
                    if not any(row):
                        continue
                    row_dict = {}
                    for i, val in enumerate(row):
                        if headers and i < len(headers) and headers[i]:
                            row_dict[headers[i]] = str(val).strip() if val is not None else ""
                    rows.append(row_dict)
        else:
            return [], {"total": 0, "valid": 0, "invalid": 0, "existing": 0}, ["Unsupported file format. Please upload a .csv or .xlsx file."]
    except Exception as e:
        return [], {"total": 0, "valid": 0, "invalid": 0, "existing": 0}, [f"Error reading file: {e}"]

    parsed_users = []
    valid_count = 0
    invalid_count = 0
    existing_count = 0

    for idx, row in enumerate(rows, start=2):
        n_row = {str(k).strip().lower(): str(v).strip() for k, v in row.items() if k is not None}
        username = n_row.get("username", "").strip()
        email = n_row.get("email", "").strip()
        display_name = n_row.get("display_name", n_row.get("displayname", n_row.get("name", username))).strip() or username
        password = n_row.get("password", "").strip()
        role = n_row.get("role", "contributor").strip().lower()
        if role not in ["admin", "contributor", "reader"]:
            role = "contributor"

        enforce_reset_val = n_row.get("enforcepasswordreset", n_row.get("enforce_reset", n_row.get("must_change_password", "yes"))).strip().lower()
        must_change = enforce_reset_val in ["yes", "true", "1", "y"]

        raw_expiry = n_row.get("password_expiry_days", n_row.get("expiry_days", n_row.get("expiry", "")))
        try:
            password_expiry_days = int(raw_expiry) if raw_expiry else DEFAULT_PASSWORD_EXPIRY_DAYS
        except Exception:
            password_expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

        row_errors = []
        if not username:
            row_errors.append("Missing username")
        if not password:
            row_errors.append(f"Missing password for '{username or 'user'}'")
        elif len(password) < 6:
            row_errors.append(f"Password must be at least 6 characters for '{username}'")

        is_valid = len(row_errors) == 0
        already_exists = False
        if username and client:
            try:
                if get_user_by_username(client, username) is not None:
                    already_exists = True
                    existing_count += 1
            except Exception:
                pass

        if is_valid:
            valid_count += 1
        else:
            invalid_count += 1
            errors.append(f"Row {idx}: {', '.join(row_errors)}")

        user_item = {
            "row_index": idx,
            "username": username,
            "display_name": display_name,
            "email": email,
            "password": password,
            "masked_password": "••••••••" if password else "(empty)",
            "role": role,
            "must_change_password": must_change,
            "password_expiry_days": password_expiry_days,
            "is_valid": is_valid,
            "already_exists": already_exists,
            "error_msg": ", ".join(row_errors) if row_errors else ""
        }
        parsed_users.append(user_item)

    summary = {
        "total": len(parsed_users),
        "valid": valid_count,
        "invalid": invalid_count,
        "existing": existing_count
    }
    return parsed_users, summary, errors

def bulk_create_users_from_file(client, file_obj, filename):
    """Bulk imports users from CSV or Excel file."""
    parsed_users, summary, errors = parse_users_file(client, file_obj, filename)
    created = 0
    skipped = 0

    for u in parsed_users:
        if not u["is_valid"]:
            skipped += 1
            continue
        try:
            save_user(
                client=client,
                username=u["username"],
                password=u["password"],
                email=u["email"],
                display_name=u["display_name"],
                role=u["role"],
                is_active=True,
                must_change_password=u["must_change_password"],
                password_expiry_days=u["password_expiry_days"]
            )
            created += 1
        except Exception as e:
            skipped += 1
            errors.append(f"Row {u['row_index']} ({u['username']}): {e}")

    return created, skipped, errors

# ---------- Activity Audit Logging Helper ----------

def log_activity(service, action, target, status='SUCCESS', details='', username=None, role=None):
    """Writes an immutable activity audit log entry into LOGS_CONTAINER in Cosmos DB."""
    try:
        client = get_cosmos_client()
        if not client:
            return
        container = get_system_logs_container(client)

        now = datetime.now(timezone.utc)
        inv_ts = f"{9999999999 - int(now.timestamp()):010d}"
        log_id = f"{inv_ts}_{uuid.uuid4().hex[:8]}"

        user_info = {}
        ip = "127.0.0.1"
        if has_request_context():
            try:
                user_info = session.get("user", {}) if session else {}
            except Exception:
                user_info = {}
            try:
                ip = request.headers.get("X-Forwarded-For", request.remote_addr or "127.0.0.1").split(",")[0].strip()
            except Exception:
                ip = "127.0.0.1"

        uname = username or user_info.get("username") or "Anonymous"
        urole = role or user_info.get("role") or "N/A"

        log_doc = {
            "id": log_id,
            "timestamp": now.isoformat(),
            "timestamp_formatted": now.strftime("%Y-%m-%d %H:%M:%S UTC"),
            "username": uname,
            "role": urole,
            "service": service,
            "action": action,
            "target": str(target)[:500] if target else "",
            "status": status,
            "details": str(details)[:1000] if details else "",
            "ip_address": ip
        }
        try:
            execute_with_429_retry(container.create_item, body=log_doc)
        except cosmos_exceptions.CosmosResourceNotFoundError:
            container = ensure_system_container(client, LOGS_CONTAINER, LOGS_PK_PATH)
            execute_with_429_retry(container.create_item, body=log_doc)
    except Exception as e:
        print(f"Activity logging error: {e}")

def cleanup_old_activity_logs(client=None, retention_days=None):
    """Purges activity logs older than retention_days (default ACTIVITY_LOG_RETENTION_DAYS=30 days) from LOGS_CONTAINER."""
    try:
        cl = client or get_cosmos_client()
        if not cl:
            return 0
        days = retention_days if retention_days is not None else ACTIVITY_LOG_RETENTION_DAYS
        container = get_system_logs_container(cl)
        cutoff_dt = datetime.now(timezone.utc) - timedelta(days=days)
        cutoff_iso = cutoff_dt.isoformat()

        pk_path = normalize_pk_path(LOGS_PK_PATH)
        pk_key = pk_path.strip("/")

        old_items = list(container.query_items(
            query="SELECT c.id, c.timestamp FROM c WHERE c.timestamp < @cutoff",
            parameters=[{"name": "@cutoff", "value": cutoff_iso}],
            enable_cross_partition_query=True
        ))
        deleted_count = 0
        for item in old_items:
            try:
                item_id = item["id"]
                pk_val = item.get(pk_key, item_id)
                execute_with_429_retry(container.delete_item, item=item_id, partition_key=pk_val)
                deleted_count += 1
            except Exception:
                pass
        return deleted_count
    except Exception as e:
        print(f"Error cleaning up old activity logs: {e}")
        return 0

def query_activity_logs(client, service=None, username=None, status=None, from_date=None, to_date=None, limit=250):
    """Queries activity audit logs from LOGS_CONTAINER with filtering and 30-day retention enforcement."""
    try:
        container = get_system_logs_container(client)

        retention_cutoff_dt = (datetime.now(timezone.utc) - timedelta(days=ACTIVITY_LOG_RETENTION_DAYS)).date()
        retention_cutoff_iso = retention_cutoff_dt.isoformat() + "T00:00:00"

        query = "SELECT * FROM c WHERE c.timestamp >= @retention_cutoff"
        params = [{"name": "@retention_cutoff", "value": retention_cutoff_iso}]

        if service and service != "all":
            query += " AND LOWER(c.service) = @service"
            params.append({"name": "@service", "value": service.lower().strip()})
        if username and username.strip():
            query += " AND CONTAINS(LOWER(c.username), @uname)"
            params.append({"name": "@uname", "value": username.lower().strip()})
        if status and status != "all":
            query += " AND c.status = @status"
            params.append({"name": "@status", "value": status.strip()})
        if from_date and from_date.strip():
            try:
                parsed_from = datetime.strptime(from_date.strip(), "%Y-%m-%d").date()
                effective_from = max(parsed_from, retention_cutoff_dt)
            except Exception:
                effective_from = retention_cutoff_dt
            query += " AND c.timestamp >= @from_date"
            params.append({"name": "@from_date", "value": f"{effective_from.isoformat()}T00:00:00"})
        if to_date and to_date.strip():
            query += " AND c.timestamp <= @to_date"
            params.append({"name": "@to_date", "value": f"{to_date.strip()}T23:59:59"})

        query += f" ORDER BY c.id ASC OFFSET 0 LIMIT {limit}"
        return list(container.query_items(query=query, parameters=params, enable_cross_partition_query=True))
    except cosmos_exceptions.CosmosResourceNotFoundError:
        return []
    except Exception as e:
        print(f"Error querying activity logs: {e}")
        return []

# ---------- Helper Functions ----------
def execute_with_429_retry(func, *args, max_retries=10, initial_delay=0.1, **kwargs):
    """
    Executes a Cosmos DB SDK operation with automatic HTTP 429 exponential backoff retry.
    Returns: (result, retries_count)
    """
    delay = initial_delay
    retries = 0
    for attempt in range(max_retries):
        try:
            return func(*args, **kwargs), retries
        except cosmos_exceptions.CosmosHttpResponseError as e:
            if e.status_code == 429: # RequestRateTooLarge
                retries += 1
                retry_after_ms = e.headers.get("x-ms-retry-after-ms") if hasattr(e, "headers") and e.headers else None
                if retry_after_ms:
                    try:
                        sleep_time = (float(retry_after_ms) / 1000.0) + random.uniform(0.01, 0.05)
                    except Exception:
                        sleep_time = delay + random.uniform(0.01, 0.05)
                else:
                    sleep_time = delay + random.uniform(0.01, 0.05)
                    delay = min(delay * 2, 5.0)
                time.sleep(sleep_time)
            else:
                raise e
    # Final attempt
    return func(*args, **kwargs), retries

def get_partition_key_path(container):
    """Programmatically fetch the partition key path for a container."""
    try:
        properties = container.read()
        paths = properties.get("partitionKey", {}).get("paths", [])
        return paths[0] if paths else None
    except Exception as e:
        print(f"Error fetching partition key path: {e}")
        return "/id" # fallback

def extract_cosmos_error_message(err):
    """Extract a readable, clean message from Azure Cosmos DB SDK exceptions."""
    err_str = str(err)
    try:
        import re
        json_matches = re.findall(r'(\{[^{}]*"errors"[^{}]*\})', err_str, re.DOTALL)
        if not json_matches:
            json_matches = re.findall(r'(\{.*"errors".*\})', err_str, re.DOTALL)
        for jm in json_matches:
            try:
                data = json.loads(jm)
                errors = data.get("errors", [])
                if errors and isinstance(errors, list):
                    first = errors[0]
                    code = first.get("code", "")
                    msg = first.get("message", "")
                    if code and msg:
                        return f"[{code}] {msg}"
                    elif msg:
                        return msg
            except Exception:
                continue

        msg_match = re.search(r'"message":\s*"([^"]+)"', err_str)
        if msg_match:
            code_match = re.search(r'"code":\s*"([^"]+)"', err_str)
            if code_match:
                return f"[{code_match.group(1)}] {msg_match.group(1)}"
            return msg_match.group(1)
    except Exception:
        pass

    if "Content:" in err_str:
        err_str = err_str.split("Content:")[0].strip()
    return err_str.strip()

def extract_partition_key_value(item, pk_path):
    """Dynamically extract the partition key value from an item's attributes based on the path."""
    if not pk_path:
        return None
    parts = pk_path.strip("/").split("/")
    val = item
    for p in parts:
        if isinstance(val, dict) and p in val:
            val = val[p]
        else:
            return None
    return val

def parse_and_build_query(search_mode, search_query, offset=0, limit=10):
    """
    Processes user search input into valid Cosmos DB SQL.
    Supports:
    1. Simple ID search (CONTAINS)
    2. Shorthand WHERE clauses (e.g. c.status = 'active')
    3. Full Cosmos SQL queries (e.g. SELECT c.id, c.name FROM c WHERE ... ORDER BY ...)
    4. Complex queries (GROUP BY, VALUE COUNT(1), TOP ...)
    """
    if not search_query or not search_query.strip():
        return {
            "items_sql": f"SELECT * FROM c ORDER BY c.id OFFSET {offset} LIMIT {limit}",
            "count_sql": "SELECT VALUE COUNT(1) FROM c",
            "params": [],
            "is_custom_projection": False,
            "is_direct_query": False,
            "can_paginate": True
        }

    query = search_query.strip()

    if search_mode == "simple":
        where = "WHERE CONTAINS(LOWER(c.id), @search)"
        params = [{"name": "@search", "value": query.lower()}]
        return {
            "items_sql": f"SELECT * FROM c {where} ORDER BY c.id OFFSET {offset} LIMIT {limit}",
            "count_sql": f"SELECT VALUE COUNT(1) FROM c {where}",
            "params": params,
            "is_custom_projection": False,
            "is_direct_query": False,
            "can_paginate": True
        }

    # Advanced Mode
    upper_q = query.upper()

    # Check if it's a full SELECT query
    if upper_q.startswith("SELECT"):
        has_group_by = "GROUP BY" in upper_q
        has_value = "SELECT VALUE" in upper_q
        has_top = "SELECT TOP" in upper_q
        has_offset_limit = "OFFSET" in upper_q and "LIMIT" in upper_q

        if has_group_by or has_value or has_top or has_offset_limit:
            return {
                "items_sql": query,
                "count_sql": None,
                "params": [],
                "is_custom_projection": not upper_q.startswith("SELECT * FROM C"),
                "is_direct_query": True,
                "can_paginate": False
            }

        # For general SELECT ... FROM c [WHERE ...] [ORDER BY ...]
        count_sql = None
        from_idx = upper_q.find("FROM")
        if from_idx != -1:
            order_idx = upper_q.rfind("ORDER BY")
            if order_idx != -1 and order_idx > from_idx:
                from_where_part = query[from_idx:order_idx].strip()
            else:
                from_where_part = query[from_idx:].strip()
            count_sql = f"SELECT VALUE COUNT(1) {from_where_part}"

        if "ORDER BY" in upper_q:
            items_sql = f"{query} OFFSET {offset} LIMIT {limit}"
        else:
            items_sql = f"{query} ORDER BY c.id OFFSET {offset} LIMIT {limit}"

        return {
            "items_sql": items_sql,
            "count_sql": count_sql,
            "params": [],
            "is_custom_projection": not (upper_q.startswith("SELECT * FROM C") or upper_q.startswith("SELECT * FROM C ")),
            "is_direct_query": False,
            "can_paginate": True
        }
    else:
        # Shorthand WHERE clause
        if upper_q.startswith("WHERE"):
            clean_where = query[5:].strip()
        else:
            clean_where = query.strip()

        return {
            "items_sql": f"SELECT * FROM c WHERE {clean_where} ORDER BY c.id OFFSET {offset} LIMIT {limit}",
            "count_sql": f"SELECT VALUE COUNT(1) FROM c WHERE {clean_where}",
            "params": [],
            "is_custom_projection": False,
            "is_direct_query": False,
            "can_paginate": True
        }

# Legacy helper for export compatibility
def build_search_query(search_mode, search_query):
    meta = parse_and_build_query(search_mode, search_query)
    if search_mode == "advanced":
        if search_query.strip().upper().startswith("SELECT"):
            return "", []
        clean = search_query.strip()
        if clean.upper().startswith("WHERE"):
            clean = clean[5:].strip()
        return f"WHERE {clean}" if clean else "", []
    elif search_mode == "simple" and search_query.strip():
        return "WHERE CONTAINS(LOWER(c.id), @search)", [{"name": "@search", "value": search_query.strip().lower()}]
    return "", []

# ---------- Root Redirect ----------
@app.route("/")
def index_redirect():
    return redirect(url_for("ui.dashboard"))

# ---------- Authentication & Portal Setup Routes ----------

@ui.route("/login", methods=["GET", "POST"])
def login():
    auto_connect_from_env_if_available()
    cosmos_connected = is_cosmos_connected()
    client = get_cosmos_client()

    if cosmos_connected and client:
        ensure_system_containers_exist(client)
        if not check_has_any_users(client):
            return redirect(url_for("ui.setup"))

    if request.method == "POST":
        login_type = request.form.get("login_type")

        # 1. User Authentication (Username + Password)
        if login_type == "user_auth" or ("username" in request.form and "password" in request.form and not request.form.get("conn_string")):
            if not cosmos_connected or not client:
                flash("Cosmos DB backend is not connected. Please connect first.", "warning")
                return redirect(url_for("ui.login"))

            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")

            user = get_user_by_username(client, username)
            if not user or not check_password_hash(user.get("password_hash", ""), password):
                log_activity("auth", "LOGIN_FAILED", username, status="FAILED", details="Invalid username or password", username=username, role="N/A")
                flash("Invalid username or password.", "danger")
                return redirect(url_for("ui.login"))

            if not user.get("is_active", True):
                log_activity("auth", "LOGIN_BLOCKED", username, status="FAILED", details="Account is disabled", username=username, role=user.get("role"))
                flash("Your account is disabled. Please contact your administrator.", "danger")
                return redirect(url_for("ui.login"))

            # Check if password reset is enforced
            if user.get("must_change_password", False):
                session["pending_user"] = username
                return redirect(url_for("ui.force_password_reset"))

            # Check password expiry policy
            try:
                expiry_days = int(user.get("password_expiry_days", DEFAULT_PASSWORD_EXPIRY_DAYS) or 0)
            except Exception:
                expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

            if expiry_days > 0:
                last_set = user.get("password_last_set_at") or user.get("created_at")
                if last_set:
                    try:
                        last_set_dt = datetime.fromisoformat(last_set.replace("Z", "+00:00"))
                        age_days = (datetime.now(timezone.utc) - last_set_dt).days
                        if age_days >= expiry_days:
                            save_user(client, username=username, must_change_password=True)
                            session["pending_user"] = username
                            flash(f"Your password has expired ({age_days} days old, max policy {expiry_days} days). Please create a new password.", "warning")
                            return redirect(url_for("ui.force_password_reset"))
                    except Exception as e:
                        print(f"Password expiry check error: {e}")

            # Update last login & establish authenticated session
            save_user(client, username=username, update_login=True)
            session["user"] = {
                "username": user["username"],
                "display_name": user.get("display_name") or user["username"],
                "email": user.get("email", ""),
                "role": user.get("role", "contributor"),
                "is_authenticated": True
            }
            log_activity("auth", "LOGIN", username, status="SUCCESS", username=username, role=user.get("role"))
            flash(f"Welcome back, {session['user']['display_name']}!", "success")
            return redirect(url_for("ui.dashboard"))

        else:
            # 2. Cosmos DB Backend Connection Submitted
            return connect_cosmos()

    if require_auth():
        return redirect(url_for("ui.dashboard"))

    return render_template("login.html", cosmos_connected=cosmos_connected)


@ui.route("/connect-cosmos", methods=["POST"])
def connect_cosmos():
    auth_method = request.form.get("auth_method")
    endpoint = request.form.get("endpoint", "").strip()
    conn_string = request.form.get("conn_string", "").strip()
    account_key = request.form.get("account_key", "").strip()
    tenant_id = request.form.get("tenant_id", "").strip()
    client_id = request.form.get("client_id", "").strip()
    client_secret = request.form.get("client_secret", "").strip()

    try:
        client = None
        auth_info = {}

        if auth_method == "conn_str":
            if not conn_string:
                raise ValueError("Connection string is required.")
            client = CosmosClient.from_connection_string(conn_string)
            for part in conn_string.split(";"):
                if part.startswith("AccountEndpoint="):
                    auth_info["endpoint"] = part.replace("AccountEndpoint=", "")
            auth_info["method"] = "Connection String"

        elif auth_method == "key":
            if not endpoint or not account_key:
                raise ValueError("Endpoint URI and Account Key are required.")
            client = CosmosClient(endpoint, credential=account_key, connection_verify=True)
            auth_info["endpoint"] = endpoint
            auth_info["method"] = "Account Key"

        elif auth_method == "sp":
            if not endpoint or not tenant_id or not client_id or not client_secret:
                raise ValueError("All Service Principal fields are required.")
            credential = ClientSecretCredential(tenant_id, client_id, client_secret)
            client = CosmosClient(endpoint, credential=credential, connection_verify=True)
            auth_info["endpoint"] = endpoint
            auth_info["method"] = "Service Principal"
        else:
            raise ValueError("Invalid authentication method selected.")

        # Test connection by listing databases
        _ = list(client.list_databases())

        sid = str(uuid.uuid4())
        session["sid"] = sid
        session["auth_info"] = auth_info
        CLIENT_STORE[sid] = {
            "client": client,
            "login_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }

        ensure_system_containers_exist(client)
        if not check_has_any_users(client):
            return redirect(url_for("ui.setup"))

        flash("Connected to Cosmos DB. Please sign in.", "success")
        return redirect(url_for("ui.login"))

    except Exception as e:
        traceback.print_exc()
        flash(f"Connection failed: {str(e)}", "danger")
        return redirect(url_for("ui.login"))


@ui.route("/setup", methods=["GET", "POST"])
def setup():
    auto_connect_from_env_if_available()
    if not is_cosmos_connected():
        flash("Please configure Cosmos DB connection first.", "warning")
        return redirect(url_for("ui.login"))

    client = get_cosmos_client()
    ensure_system_containers_exist(client)
    if check_has_any_users(client):
        flash("Setup has already been completed.", "info")
        return redirect(url_for("ui.login"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        display_name = request.form.get("display_name", "").strip() or "System Administrator"
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not username or not password:
            flash("Username and password are required.", "danger")
            return render_template("setup.html")
        if password != confirm_password:
            flash("Passwords do not match.", "danger")
            return render_template("setup.html")
        if len(password) < 6:
            flash("Password must be at least 6 characters.", "danger")
            return render_template("setup.html")

        try:
            save_user(
                client=client,
                username=username,
                password=password,
                email=email,
                display_name=display_name,
                role="admin",
                is_active=True,
                must_change_password=False,
                password_expiry_days=0,
                update_login=True
            )
            session["user"] = {
                "username": username,
                "display_name": display_name,
                "email": email,
                "role": "admin",
                "is_authenticated": True
            }
            log_activity("auth", "INITIAL_SETUP", username, status="SUCCESS", details="Master Admin Created", username=username, role="admin")
            flash("Master administrator account initialized successfully! Welcome to Cosmos UI.", "success")
            return redirect(url_for("ui.dashboard"))
        except Exception as e:
            traceback.print_exc()
            err_str = str(e)
            if "Owner resource does not exist" in err_str or "NotFound" in err_str:
                flash(
                    f"Cosmos DB system database '{SYSTEM_DB}' or container '{USER_CONTAINER}' does not exist, and your Service Principal does not have permission to create databases/containers. "
                    f"Please either: (1) Manually create Database '{SYSTEM_DB}' and Containers '{USER_CONTAINER}' (PK: {USER_PK_PATH}) and '{LOGS_CONTAINER}' (PK: {LOGS_PK_PATH}) in Azure Portal, or "
                    f"(2) Set environment variable COSMOS_AUTH_DB to an existing database your Service Principal can access.",
                    "danger"
                )
            else:
                flash(f"Failed to create admin account: {err_str}", "danger")
            return render_template("setup.html")

    return render_template("setup.html")


@ui.route("/force-password-reset", methods=["GET", "POST"])
def force_password_reset():
    username = session.get("pending_user")
    if not username:
        return redirect(url_for("ui.login"))

    client = get_cosmos_client()
    if request.method == "POST":
        new_password = request.form.get("new_password", "")
        confirm_password = request.form.get("confirm_password", "")

        if len(new_password) < 6:
            flash("New password must be at least 6 characters.", "danger")
            return render_template("force_password_reset.html")
        if new_password != confirm_password:
            flash("Passwords do not match.", "danger")
            return render_template("force_password_reset.html")

        try:
            save_user(client=client, username=username, password=new_password, must_change_password=False, update_login=True)
            user = get_user_by_username(client, username)
            session.pop("pending_user", None)
            session["user"] = {
                "username": user["username"],
                "display_name": user.get("display_name") or user["username"],
                "email": user.get("email", ""),
                "role": user.get("role", "contributor"),
                "is_authenticated": True
            }
            log_activity("auth", "PASSWORD_RESET_FORCED", username, status="SUCCESS", username=username, role=user.get("role"))
            flash("Password updated successfully! Welcome to Cosmos UI.", "success")
            return redirect(url_for("ui.dashboard"))
        except Exception as e:
            flash(f"Error resetting password: {str(e)}", "danger")
            return render_template("force_password_reset.html")

    return render_template("force_password_reset.html")


@ui.route("/change-my-password", methods=["POST"])
@login_required
def change_my_password():
    client = get_cosmos_client()
    username = session["user"]["username"]
    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")

    if len(new_password) < 6:
        flash("New password must be at least 6 characters.", "danger")
        return redirect(request.referrer or url_for("ui.dashboard"))
    if new_password != confirm_password:
        flash("New passwords do not match.", "danger")
        return redirect(request.referrer or url_for("ui.dashboard"))

    user = get_user_by_username(client, username)
    if not user or not check_password_hash(user.get("password_hash", ""), current_password):
        flash("Current password incorrect.", "danger")
        return redirect(request.referrer or url_for("ui.dashboard"))

    try:
        save_user(client=client, username=username, password=new_password, must_change_password=False)
        log_activity("auth", "PASSWORD_CHANGE", username, status="SUCCESS", username=username, role=session["user"]["role"])
        flash("Your password has been changed successfully.", "success")
    except Exception as e:
        flash(f"Error changing password: {str(e)}", "danger")

    return redirect(request.referrer or url_for("ui.dashboard"))


@ui.route("/logout")
def logout():
    uname = session.get("user", {}).get("username", "Anonymous")
    log_activity("auth", "LOGOUT", uname, status="SUCCESS")
    session.pop("user", None)
    flash("Signed out successfully.", "info")
    return redirect(url_for("ui.login"))


# ---------- Portal Management: User & Role Management (Admin Only) ----------

@ui.route("/users")
@admin_required
def users_list():
    client = get_cosmos_client()
    users = get_all_users(client)
    
    # Load sidebar tree excluding system db
    databases = [db for db in client.list_databases() if db["id"] != SYSTEM_DB]
    db_tree = []
    for db in databases:
        dbc = client.get_database_client(db["id"])
        db_tree.append({
            "id": db["id"],
            "containers": [c["id"] for c in dbc.list_containers()]
        })

    return render_template("users.html", users=users, db_tree=db_tree, auth_info=session.get("auth_info"))


@ui.route("/users/create", methods=["POST"])
@admin_required
def create_user():
    client = get_cosmos_client()
    username = request.form.get("username", "").strip()
    display_name = request.form.get("display_name", "").strip()
    email = request.form.get("email", "").strip()
    role = request.form.get("role", "contributor").strip().lower()
    password = request.form.get("password", "")
    enforce_reset = bool(request.form.get("enforce_reset"))
    expiry_days_raw = request.form.get("password_expiry_days")
    try:
        expiry_days = int(expiry_days_raw) if expiry_days_raw is not None and expiry_days_raw != "" else DEFAULT_PASSWORD_EXPIRY_DAYS
    except (ValueError, TypeError):
        expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

    if not username or not password:
        flash("Username and password are required.", "danger")
        return redirect(url_for("ui.users_list"))

    if len(password) < 6:
        flash("Password must be at least 6 characters.", "danger")
        return redirect(url_for("ui.users_list"))

    existing = get_user_by_username(client, username)
    if existing:
        flash(f"User '{username}' already exists.", "warning")
        return redirect(url_for("ui.users_list"))

    try:
        save_user(
            client=client,
            username=username,
            password=password,
            email=email,
            display_name=display_name,
            role=role,
            is_active=True,
            must_change_password=enforce_reset,
            password_expiry_days=expiry_days
        )
        log_activity("user_mgmt", "CREATE_USER", username, status="SUCCESS", details=f"Role: {role}, Email: {email}, Expiry: {expiry_days}d")
        flash(f"User '{username}' created successfully.", "success")
    except Exception as e:
        log_activity("user_mgmt", "CREATE_USER", username, status="FAILED", details=str(e))
        flash(f"Failed to create user: {str(e)}", "danger")

    return redirect(url_for("ui.users_list"))


@ui.route("/users/bulk-preview", methods=["POST"])
@admin_required
def bulk_preview_users():
    client = get_cosmos_client()
    uploaded_file = request.files.get("file")
    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"success": False, "error": "No file was uploaded."}), 400

    parsed_users, summary, errors = parse_users_file(client, uploaded_file, uploaded_file.filename)
    if not parsed_users and errors:
        return jsonify({"success": False, "error": "; ".join(errors)}), 400

    # Store raw parsed list in session for one-click confirmation
    session["bulk_preview_users"] = parsed_users

    # Return sanitized view payload with masked passwords
    preview_users = []
    for u in parsed_users:
        preview_users.append({
            "row_index": u["row_index"],
            "username": u["username"],
            "display_name": u["display_name"],
            "email": u["email"],
            "role": u["role"],
            "masked_password": u["masked_password"],
            "must_change_password": u["must_change_password"],
            "password_expiry_days": u["password_expiry_days"],
            "is_valid": u["is_valid"],
            "already_exists": u["already_exists"],
            "error_msg": u["error_msg"]
        })

    return jsonify({
        "success": True,
        "summary": summary,
        "users": preview_users,
        "errors": errors
    })


@ui.route("/users/bulk-create", methods=["POST"])
@admin_required
def bulk_create_users():
    client = get_cosmos_client()
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest" or request.is_json
    data = request.get_json(silent=True) or {}
    confirm = data.get("confirm") or request.form.get("confirm")

    if confirm:
        cached_users = session.pop("bulk_preview_users", None)
        if not cached_users:
            if is_ajax:
                return jsonify({"success": False, "error": "Preview session expired. Please re-upload the file."}), 400
            flash("Preview session expired. Please re-upload the file.", "warning")
            return redirect(url_for("ui.users_list"))

        created = 0
        skipped = 0
        errors = []
        for u in cached_users:
            if not u["is_valid"]:
                skipped += 1
                continue
            try:
                save_user(
                    client=client,
                    username=u["username"],
                    password=u["password"],
                    email=u["email"],
                    display_name=u["display_name"],
                    role=u["role"],
                    is_active=True,
                    must_change_password=u["must_change_password"],
                    password_expiry_days=u.get("password_expiry_days", DEFAULT_PASSWORD_EXPIRY_DAYS)
                )
                created += 1
            except Exception as e:
                skipped += 1
                errors.append(f"Row {u['row_index']} ({u['username']}): {e}")

        details = f"Created: {created}, Skipped: {skipped}"
        if errors:
            details += f" (Errors: {'; '.join(errors[:3])})"
        log_activity("user_mgmt", "BULK_CREATE_USERS", "Bulk Import (Preview Confirmed)", status="SUCCESS" if created > 0 else "FAILED", details=details)

        msg = f"Bulk import complete: {created} user(s) created."
        if skipped > 0:
            msg += f" {skipped} skipped."
        if is_ajax:
            return jsonify({"success": True, "created": created, "skipped": skipped, "errors": errors, "message": msg})
        flash(msg, "success" if created > 0 else "warning")
        return redirect(url_for("ui.users_list"))

    # Direct file upload fallback
    uploaded_file = request.files.get("file")
    if not uploaded_file or not uploaded_file.filename:
        if is_ajax:
            return jsonify({"success": False, "error": "No file was uploaded."}), 400
        flash("No file was uploaded.", "warning")
        return redirect(url_for("ui.users_list"))

    created, skipped, errors = bulk_create_users_from_file(client, uploaded_file, uploaded_file.filename)
    details = f"Created: {created}, Skipped: {skipped}"
    if errors:
        details += f" (Errors: {'; '.join(errors[:3])})"
    log_activity("user_mgmt", "BULK_CREATE_USERS", uploaded_file.filename, status="SUCCESS" if created > 0 else "FAILED", details=details)

    msg = f"Bulk import complete: {created} user(s) created."
    if skipped > 0:
        msg += f" {skipped} skipped. {'; '.join(errors)}"
    if is_ajax:
        return jsonify({"success": True, "created": created, "skipped": skipped, "errors": errors, "message": msg})
    flash(msg, "success" if created > 0 else "warning")
    return redirect(url_for("ui.users_list"))


@ui.route("/users/edit", methods=["POST"])
@admin_required
def edit_user():
    client = get_cosmos_client()
    username = request.form.get("username", "").strip()
    display_name = request.form.get("display_name", "").strip()
    email = request.form.get("email", "").strip()
    role = request.form.get("role", "contributor").strip().lower()
    expiry_days_raw = request.form.get("password_expiry_days")
    try:
        expiry_days = int(expiry_days_raw) if expiry_days_raw is not None and expiry_days_raw != "" else DEFAULT_PASSWORD_EXPIRY_DAYS
    except (ValueError, TypeError):
        expiry_days = DEFAULT_PASSWORD_EXPIRY_DAYS

    try:
        save_user(
            client=client,
            username=username,
            display_name=display_name,
            email=email,
            role=role,
            password_expiry_days=expiry_days
        )
        if session.get("user", {}).get("username") == username:
            session["user"]["display_name"] = display_name or username
            session["user"]["email"] = email
            session["user"]["role"] = role

        log_activity("user_mgmt", "EDIT_USER", username, status="SUCCESS", details=f"Role: {role}, Email: {email}, Expiry: {expiry_days}d")
        flash(f"User '{username}' updated successfully.", "success")
    except Exception as e:
        log_activity("user_mgmt", "EDIT_USER", username, status="FAILED", details=str(e))
        flash(f"Failed to update user: {str(e)}", "danger")

    return redirect(url_for("ui.users_list"))


@ui.route("/users/reset-password", methods=["POST"])
@admin_required
def reset_user_password():
    client = get_cosmos_client()
    username = request.form.get("username", "").strip()
    new_password = request.form.get("new_password", "")
    enforce_reset = bool(request.form.get("enforce_reset"))

    if len(new_password) < 6:
        flash("Password must be at least 6 characters.", "danger")
        return redirect(url_for("ui.users_list"))

    try:
        save_user(client=client, username=username, password=new_password, must_change_password=enforce_reset)
        log_activity("user_mgmt", "RESET_PASSWORD", username, status="SUCCESS", details=f"Enforce Reset: {enforce_reset}")
        flash(f"Password for '{username}' has been reset.", "success")
    except Exception as e:
        log_activity("user_mgmt", "RESET_PASSWORD", username, status="FAILED", details=str(e))
        flash(f"Failed to reset password: {str(e)}", "danger")

    return redirect(url_for("ui.users_list"))


@ui.route("/users/toggle-status", methods=["POST"])
@admin_required
def toggle_user_status():
    client = get_cosmos_client()
    username = request.form.get("username", "").strip()
    is_active_val = request.form.get("is_active", "1") == "1"

    if username == session.get("user", {}).get("username") and not is_active_val:
        flash("You cannot disable your own administrator account.", "warning")
        return redirect(url_for("ui.users_list"))

    try:
        save_user(client=client, username=username, is_active=is_active_val)
        status_name = "enabled" if is_active_val else "disabled"
        log_activity("user_mgmt", f"{'ENABLE' if is_active_val else 'DISABLE'}_USER", username, status="SUCCESS")
        flash(f"Account for '{username}' has been {status_name}.", "success")
    except Exception as e:
        log_activity("user_mgmt", "TOGGLE_STATUS_USER", username, status="FAILED", details=str(e))
        flash(f"Failed to update status: {str(e)}", "danger")

    return redirect(url_for("ui.users_list"))


@ui.route("/users/delete", methods=["POST"])
@admin_required
def delete_user():
    client = get_cosmos_client()
    username = request.form.get("username", "").strip()

    if username == session.get("user", {}).get("username"):
        flash("You cannot delete your own administrator account.", "warning")
        return redirect(url_for("ui.users_list"))

    try:
        delete_user_by_username(client, username)
        log_activity("user_mgmt", "DELETE_USER", username, status="SUCCESS")
        flash(f"User '{username}' deleted successfully.", "success")
    except Exception as e:
        log_activity("user_mgmt", "DELETE_USER", username, status="FAILED", details=str(e))
        flash(f"Failed to delete user: {str(e)}", "danger")

    return redirect(url_for("ui.users_list"))


@ui.route("/users/template.csv")
@admin_required
def download_user_template():
    csv_data = "username,email,display_name,password,enforcepasswordreset,role,password_expiry_days\n"
    csv_data += "john,john@example.com,John Doe,Password123!,yes,contributor,90\n"
    csv_data += "sarah,sarah@example.com,Sarah Smith,TempPass456!,yes,reader,90\n"
    csv_data += "admin2,admin2@example.com,Second Admin,SecureAdmin789!,no,admin,0\n"
    return Response(
        csv_data,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=users_import_template.csv"}
    )


# ---------- Portal Management: Activity Logs (Admin Only) ----------

@ui.route("/activity-logs")
@admin_required
def activity_logs():
    client = get_cosmos_client()
    filter_service = request.args.get("service", "all")
    filter_username = request.args.get("username", "")
    filter_status = request.args.get("status", "all")
    from_date = request.args.get("from_date", "")
    to_date = request.args.get("to_date", "")

    # Trigger asynchronous background cleanup of logs older than 30 days
    try:
        threading.Thread(target=cleanup_old_activity_logs, args=(client,), daemon=True).start()
    except Exception:
        pass

    logs = query_activity_logs(
        client=client,
        service=filter_service,
        username=filter_username,
        status=filter_status,
        from_date=from_date,
        to_date=to_date,
        limit=250
    )

    # Load sidebar tree excluding system db
    databases = [db for db in client.list_databases() if db["id"] != SYSTEM_DB]
    db_tree = []
    for db in databases:
        dbc = client.get_database_client(db["id"])
        db_tree.append({
            "id": db["id"],
            "containers": [c["id"] for c in dbc.list_containers()]
        })

    return render_template(
        "activity_logs.html",
        logs=logs,
        filter_service=filter_service,
        filter_username=filter_username,
        filter_status=filter_status,
        from_date=from_date,
        to_date=to_date,
        db_tree=db_tree,
        auth_info=session.get("auth_info")
    )


@ui.route("/activity-logs/export")
@admin_required
def export_activity_logs():
    client = get_cosmos_client()
    fmt = request.args.get("format", "csv").lower()
    filter_service = request.args.get("service", "all")
    filter_username = request.args.get("username", "")
    filter_status = request.args.get("status", "all")
    from_date = request.args.get("from_date", "")
    to_date = request.args.get("to_date", "")

    logs = query_activity_logs(
        client=client,
        service=filter_service,
        username=filter_username,
        status=filter_status,
        from_date=from_date,
        to_date=to_date,
        limit=2000
    )

    if fmt == "json":
        clean_logs = []
        for l in logs:
            c = dict(l)
            c.pop("_rid", None)
            c.pop("_self", None)
            c.pop("_etag", None)
            c.pop("_attachments", None)
            c.pop("_ts", None)
            clean_logs.append(c)
        return Response(
            json.dumps(clean_logs, indent=2),
            mimetype="application/json",
            headers={"Content-Disposition": "attachment;filename=cosmos_activity_logs.json"}
        )
    else: # CSV
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(["Timestamp (UTC)", "Username", "Role", "Service", "Action", "Target", "Status", "Details", "Client IP"])
        for l in logs:
            writer.writerow([
                l.get("timestamp_formatted") or l.get("timestamp", ""),
                l.get("username", ""),
                l.get("role", ""),
                l.get("service", ""),
                l.get("action", ""),
                l.get("target", ""),
                l.get("status", ""),
                l.get("details", ""),
                l.get("ip_address", "")
            ])
        return Response(
            out.getvalue().encode("utf-8-sig"),
            mimetype="text/csv",
            headers={"Content-Disposition": "attachment;filename=cosmos_activity_logs.csv"}
        )


# ---------- Dashboard & Explorer ----------

@ui.route("/")
@ui.route("/dashboard")
@login_required
def dashboard():
    client = get_cosmos_client()
    try:
        databases = [db for db in client.list_databases() if db["id"] != SYSTEM_DB]
        db_tree = []
        for db in databases:
            db_client = client.get_database_client(db["id"])
            containers = list(db_client.list_containers())
            db_tree.append({
                "id": db["id"],
                "containers": [c["id"] for c in containers]
            })
        return render_template("dashboard.html", db_tree=db_tree, auth_info=session.get("auth_info"))
    except Exception as e:
        flash(f"Error fetching databases: {str(e)}", "danger")
        return render_template("dashboard.html", db_tree=[], auth_info=session.get("auth_info"))

# ---------- Streaming File Helper for Large Ingestions ----------
def stream_file_records(file_path, filename):
    """
    Generator yielding records from JSON, JSONL, NDJSON, CSV, or XLSX files without high memory usage.
    Yields: (row_index, document_dict)
    """
    fn = filename.lower()
    
    if fn.endswith(".jsonl") or fn.endswith(".ndjson"):
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            for idx, line in enumerate(f):
                line = line.strip()
                if not line:
                    continue
                try:
                    doc = json.loads(line)
                    if isinstance(doc, dict):
                        yield idx + 1, doc
                    else:
                        yield idx + 1, {"_raw_value": doc}
                except Exception as e:
                    yield idx + 1, {"_parse_error": str(e)}

    elif fn.endswith(".json"):
        # Check if JSON array or JSON Lines format
        first_char = ""
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            for char in f.read(1024):
                if not char.isspace():
                    first_char = char
                    break
        
        if first_char == "[":
            # Standard JSON Array
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json.load(f)
                if isinstance(data, list):
                    for idx, doc in enumerate(data):
                        if isinstance(doc, dict):
                            yield idx + 1, doc
                        else:
                            yield idx + 1, {"_raw_value": doc}
                elif isinstance(data, dict):
                    yield 1, data
        else:
            # NDJSON / JSONL
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                for idx, line in enumerate(f):
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        doc = json.loads(line)
                        if isinstance(doc, dict):
                            yield idx + 1, doc
                        else:
                            yield idx + 1, {"_raw_value": doc}
                    except Exception as e:
                        yield idx + 1, {"_parse_error": str(e)}

    elif fn.endswith(".csv"):
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            headers = next(reader, None)
            if headers:
                clean_headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(headers)]
                for idx, row in enumerate(reader):
                    doc = {}
                    for h_idx, h in enumerate(clean_headers):
                        if h_idx < len(row) and row[h_idx] is not None:
                            val = row[h_idx]
                            if isinstance(val, str) and (val.startswith("{") or val.startswith("[")):
                                try:
                                    val = json.loads(val)
                                except Exception:
                                    pass
                            doc[h] = val
                    yield idx + 2, doc

    elif fn.endswith(".xlsx"):
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if headers:
            clean_headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(headers)]
            for idx, row in enumerate(rows_iter):
                if not row or all(v is None for v in row):
                    continue
                doc = {}
                for h_idx, h in enumerate(clean_headers):
                    if h_idx < len(row) and row[h_idx] is not None:
                        val = row[h_idx]
                        if isinstance(val, str) and (val.startswith("{") or val.startswith("[")):
                            try:
                                val = json.loads(val)
                            except Exception:
                                pass
                        doc[h] = val
                yield idx + 2, doc
        wb.close()


def run_bulk_import_worker(task_id, file_path, filename, db_id, container_id, pk_path, concurrency, client):
    """
    Ultra-high-throughput continuous streaming worker pool for maximum Cosmos DB ingestion speed.
    Eliminates pipeline stalls by using a continuous thread-safe producer-consumer queue.
    """
    clean_pk = pk_path.strip("/") if pk_path else "id"
    
    with IMPORT_TASKS_LOCK:
        task = IMPORT_TASKS.get(task_id)
        if not task:
            return
        task["status"] = "in_progress"
        task["start_time"] = time.time()

    try:
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)

        # Scale concurrency up to 200 workers
        num_workers = max(10, min(concurrency, 200))
        doc_queue = queue.Queue(maxsize=num_workers * 25)
        
        stats_lock = threading.Lock()
        total_processed = 0
        total_success = 0
        total_failed = 0
        total_429_retries = 0
        errors = []
        stop_event = threading.Event()

        def worker_loop():
            nonlocal total_processed, total_success, total_failed, total_429_retries
            while not stop_event.is_set():
                try:
                    item = doc_queue.get(timeout=0.15)
                except queue.Empty:
                    continue

                if item is None:
                    doc_queue.task_done()
                    break

                row_idx, doc = item
                if "_parse_error" in doc:
                    with stats_lock:
                        total_processed += 1
                        total_failed += 1
                        if len(errors) < 25:
                            errors.append(f"Row {row_idx}: Parse error: {doc['_parse_error']}")
                    doc_queue.task_done()
                    continue

                if "id" not in doc or not str(doc["id"]).strip():
                    doc["id"] = str(uuid.uuid4())
                else:
                    doc["id"] = str(doc["id"])

                pk_val = extract_partition_key_value(doc, pk_path)
                if pk_val is None:
                    doc[clean_pk] = "imported"

                try:
                    _, retries = execute_with_429_retry(container.upsert_item, body=doc, max_retries=10)
                    with stats_lock:
                        total_processed += 1
                        total_success += 1
                        total_429_retries += retries
                except Exception as err:
                    with stats_lock:
                        total_processed += 1
                        total_failed += 1
                        if len(errors) < 25:
                            errors.append(f"Row {row_idx}: {str(err)}")

                doc_queue.task_done()

        # Start consumer worker threads
        workers = []
        for _ in range(num_workers):
            t = threading.Thread(target=worker_loop, daemon=True)
            t.start()
            workers.append(t)

        # Background metric updater for smooth UI reporting
        def metric_updater():
            while not stop_event.is_set():
                time.sleep(0.2)
                with stats_lock:
                    proc = total_processed
                    succ = total_success
                    fail = total_failed
                    r429 = total_429_retries
                    errs = list(errors)
                elapsed = max(0.1, time.time() - task["start_time"])
                speed = round(proc / elapsed, 1)
                with IMPORT_TASKS_LOCK:
                    task["processed"] = proc
                    task["successful"] = succ
                    task["failed"] = fail
                    task["retries_429"] = r429
                    task["speed_per_sec"] = speed
                    task["errors"] = errs

        metric_thread = threading.Thread(target=metric_updater, daemon=True)
        metric_thread.start()

        # Producer: Stream records from file directly into queue
        records_gen = stream_file_records(file_path, filename)
        for row_info in records_gen:
            with IMPORT_TASKS_LOCK:
                if task.get("cancel_requested"):
                    stop_event.set()
                    break
            while not stop_event.is_set():
                try:
                    doc_queue.put(row_info, timeout=0.1)
                    break
                except queue.Full:
                    with IMPORT_TASKS_LOCK:
                        if task.get("cancel_requested"):
                            stop_event.set()
                            break

        # If cancelled, drain the remaining queue immediately
        if stop_event.is_set():
            while not doc_queue.empty():
                try:
                    doc_queue.get_nowait()
                    doc_queue.task_done()
                except Exception:
                    break

        # Send termination sentinels
        for _ in range(num_workers):
            try:
                doc_queue.put_nowait(None)
            except Exception:
                pass

        # Await workers with small timeout
        for t in workers:
            t.join(timeout=0.2)

        stop_event.set()
        metric_thread.join(timeout=0.3)

        elapsed = max(0.1, time.time() - task["start_time"])
        speed = round(total_processed / elapsed, 1)

        with IMPORT_TASKS_LOCK:
            if task.get("cancel_requested"):
                task["status"] = "cancelled"
            elif task.get("status") != "cancelled":
                task["status"] = "completed"
            task["end_time"] = time.time()
            task["processed"] = total_processed
            task["successful"] = total_success
            task["failed"] = total_failed
            task["retries_429"] = total_429_retries
            task["speed_per_sec"] = speed
            task["errors"] = errors

    except Exception as e:
        traceback.print_exc()
        with IMPORT_TASKS_LOCK:
            task["status"] = "failed"
            task["error_message"] = str(e)
            task["end_time"] = time.time()
            
    finally:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
        except Exception:
            pass


def run_export_worker(task_id, file_path, filename, format_type, db_id, container_id, query_sql, query_params, client):
    """
    Background worker that streams documents from Cosmos DB directly into disk
    supporting JSON, JSONL, CSV, and XLSX with 429 rate limit backoff and live progress.
    """
    with EXPORT_TASKS_LOCK:
        task = EXPORT_TASKS.get(task_id)
        if not task:
            return
        task["status"] = "in_progress"
        task["start_time"] = time.time()

    try:
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)

        processed = 0
        total_429_retries = 0

        # Query Cosmos DB with cross-partition support and 1,000-item page streaming
        query_iterable = container.query_items(
            query=query_sql,
            parameters=query_params,
            enable_cross_partition_query=True,
            max_item_count=1000
        )
        pager = query_iterable.by_page()

        def stream_items_with_retry():
            nonlocal total_429_retries
            while True:
                with EXPORT_TASKS_LOCK:
                    if task.get("cancel_requested"):
                        return

                def fetch_page():
                    try:
                        return next(pager, None)
                    except StopIteration:
                        return None

                page, retries = execute_with_429_retry(fetch_page, max_retries=10)
                total_429_retries += retries
                if page is None:
                    break
                for item in page:
                    yield item

        # Stream directly into target file format without high RAM consumption
        if format_type in ["jsonl", "ndjson"]:
            with open(file_path, "w", encoding="utf-8") as f:
                for doc in stream_items_with_retry():
                    with EXPORT_TASKS_LOCK:
                        if task.get("cancel_requested"):
                            break
                    f.write(json.dumps(doc) + "\n")
                    processed += 1
                    if processed % 200 == 0:
                        elapsed = max(0.1, time.time() - task["start_time"])
                        with EXPORT_TASKS_LOCK:
                            task["processed"] = processed
                            task["speed_per_sec"] = round(processed / elapsed, 1)
                            task["retries_429"] = total_429_retries

        elif format_type == "json":
            with open(file_path, "w", encoding="utf-8") as f:
                f.write("[\n")
                first = True
                for doc in stream_items_with_retry():
                    with EXPORT_TASKS_LOCK:
                        if task.get("cancel_requested"):
                            break
                    if not first:
                        f.write(",\n")
                    else:
                        first = False
                    f.write(json.dumps(doc, indent=2))
                    processed += 1
                    if processed % 200 == 0:
                        elapsed = max(0.1, time.time() - task["start_time"])
                        with EXPORT_TASKS_LOCK:
                            task["processed"] = processed
                            task["speed_per_sec"] = round(processed / elapsed, 1)
                            task["retries_429"] = total_429_retries
                f.write("\n]\n")

        elif format_type == "csv":
            with open(file_path, "w", encoding="utf-8", newline="") as f:
                writer = None
                headers = []
                for doc in stream_items_with_retry():
                    with EXPORT_TASKS_LOCK:
                        if task.get("cancel_requested"):
                            break
                    if writer is None:
                        if isinstance(doc, dict):
                            headers = sorted(list(doc.keys()))
                            if "id" in headers:
                                headers.remove("id")
                                headers.insert(0, "id")
                        else:
                            headers = ["value"]
                        writer = csv.writer(f)
                        writer.writerow(headers)

                    if isinstance(doc, dict):
                        row = []
                        for h in headers:
                            val = doc.get(h, "")
                            if isinstance(val, (dict, list)):
                                val = json.dumps(val)
                            row.append(val)
                        writer.writerow(row)
                    else:
                        writer.writerow([doc])

                    processed += 1
                    if processed % 200 == 0:
                        elapsed = max(0.1, time.time() - task["start_time"])
                        with EXPORT_TASKS_LOCK:
                            task["processed"] = processed
                            task["speed_per_sec"] = round(processed / elapsed, 1)
                            task["retries_429"] = total_429_retries

        elif format_type == "xlsx":
            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title="CosmosExport")
            headers_written = False
            headers = []
            for doc in stream_items_with_retry():
                with EXPORT_TASKS_LOCK:
                    if task.get("cancel_requested"):
                        break
                if not headers_written:
                    if isinstance(doc, dict):
                        headers = sorted(list(doc.keys()))
                        if "id" in headers:
                            headers.remove("id")
                            headers.insert(0, "id")
                    else:
                        headers = ["value"]
                    ws.append(headers)
                    headers_written = True

                if isinstance(doc, dict):
                    row = []
                    for h in headers:
                        val = doc.get(h, "")
                        if isinstance(val, (dict, list)):
                            val = json.dumps(val)
                        row.append(val)
                    ws.append(row)
                else:
                    ws.append([doc])

                processed += 1
                if processed % 200 == 0:
                    elapsed = max(0.1, time.time() - task["start_time"])
                    with EXPORT_TASKS_LOCK:
                        task["processed"] = processed
                        task["speed_per_sec"] = round(processed / elapsed, 1)
                        task["retries_429"] = total_429_retries
            wb.save(file_path)

        elapsed = max(0.1, time.time() - task["start_time"])
        with EXPORT_TASKS_LOCK:
            if task.get("cancel_requested"):
                task["status"] = "cancelled"
            else:
                task["status"] = "completed"
            task["end_time"] = time.time()
            task["processed"] = processed
            task["speed_per_sec"] = round(processed / elapsed, 1)
            task["retries_429"] = total_429_retries

    except Exception as e:
        traceback.print_exc()
        with EXPORT_TASKS_LOCK:
            task["status"] = "failed"
            task["error_message"] = str(e)
            task["end_time"] = time.time()


@ui.route("/db/<db_id>/container/<container_id>")
@login_required
def container_view(db_id, container_id):
    client = get_cosmos_client()
    if not client:
        flash("Cosmos DB is not connected.", "warning")
        return redirect(url_for("ui.login"))
    
    # Query parameters
    page = request.args.get("page", 1, type=int)
    limit = request.args.get("limit", 10, type=int)
    search_mode = request.args.get("search_mode", "simple")
    search_query = request.args.get("search_query", "")
    cached_total_str = request.args.get("total_items", None)

    offset = (page - 1) * limit

    try:
        # Get DB and Container Clients
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        
        # Programmatically detect Partition Key
        pk_path = get_partition_key_path(container)
        
        # Build query parts using upgraded parser
        q_meta = parse_and_build_query(search_mode, search_query, offset=offset, limit=limit)
        
        # Determine total_items (Smart Count for 20L+ records)
        total_items = None
        if q_meta["is_direct_query"]:
            total_items = None # Direct custom queries might return arbitrary rows
        elif cached_total_str is not None and str(cached_total_str).strip() != "":
            try:
                total_items = int(cached_total_str)
            except ValueError:
                total_items = None
                
        if total_items is None and q_meta["count_sql"]:
            try:
                count_iter = container.query_items(
                    query=q_meta["count_sql"],
                    parameters=q_meta["params"],
                    enable_cross_partition_query=True
                )
                count_res = list(count_iter)
                total_items = int(count_res[0]) if count_res and count_res[0] is not None else 0
            except Exception as count_err:
                print(f"Count query notice (e.g. large dataset scan): {count_err}")
                total_items = 0

        query_error = None
        raw_items = []

        # Execute items query with graceful error handling
        try:
            items_iter = container.query_items(
                query=q_meta["items_sql"],
                parameters=q_meta["params"],
                enable_cross_partition_query=True,
                max_item_count=limit
            )
            raw_items = list(items_iter)
        except Exception as q_err:
            traceback.print_exc()
            query_error = extract_cosmos_error_message(q_err)
            raw_items = []
        
        if query_error:
            total_items = 0
        elif total_items is None or total_items == 0:
            if q_meta["is_direct_query"]:
                total_items = len(raw_items)
            elif not q_meta["count_sql"]:
                total_items = len(raw_items)

        # Process items to extract partition key value, summary, or custom projections
        processed_items = []
        custom_columns = []
        all_keys = set()

        for it in raw_items:
            if isinstance(it, dict):
                has_id = "id" in it
                pk_val = extract_partition_key_value(it, pk_path) if has_id else None
                all_keys.update(it.keys())
                processed_items.append({
                    "id": it.get("id", None),
                    "pk_val": pk_val,
                    "has_id": has_id,
                    "raw": it
                })
            else:
                processed_items.append({
                    "id": None,
                    "pk_val": None,
                    "has_id": False,
                    "raw": {"_result": it}
                })

        # If custom projection without standard id, determine custom columns to display
        is_custom_projection = q_meta["is_custom_projection"] or (len(processed_items) > 0 and not any(p["has_id"] for p in processed_items))
        if is_custom_projection and all_keys:
            custom_columns = sorted([k for k in all_keys if not k.startswith("_")])[:8]
            if not custom_columns:
                custom_columns = sorted(list(all_keys))[:8]

        # Calculate pages
        effective_total = total_items if total_items is not None else len(processed_items)
        total_pages = max(1, (effective_total + limit - 1) // limit)

        # Database tree for sidebar quick-nav excluding system db
        databases = [db for db in client.list_databases() if db["id"] != SYSTEM_DB]
        db_tree = []
        for db in databases:
            dbc = client.get_database_client(db["id"])
            db_tree.append({
                "id": db["id"],
                "containers": [c["id"] for c in dbc.list_containers()]
            })

        return render_template(
            "container.html",
            db_id=db_id,
            container_id=container_id,
            pk_path=pk_path,
            items=processed_items,
            page=page,
            limit=limit,
            search_mode=search_mode,
            search_query=search_query,
            total_items=effective_total,
            total_pages=total_pages,
            is_custom_projection=is_custom_projection,
            custom_columns=custom_columns,
            can_paginate=q_meta["can_paginate"],
            query_error=query_error,
            db_tree=db_tree,
            auth_info=session.get("auth_info")
        )
    except Exception as e:
        traceback.print_exc()
        flash(f"Error accessing container: {extract_cosmos_error_message(e)}", "danger")
        return redirect(url_for("ui.dashboard"))

# ---------- API Endpoints ----------

@ui.route("/api/db/<db_id>/container/<container_id>/item", methods=["POST"])
@write_permission_required
def api_upsert_item(db_id, container_id):
    """Create or Update an item in Cosmos DB (Upsert)"""
    client = get_cosmos_client()
    
    try:
        data = request.get_json()
        if not data:
            return jsonify({"status": "error", "message": "No data provided"}), 400
        
        if "id" not in data or not str(data["id"]).strip():
            return jsonify({"status": "error", "message": "Document must contain an 'id' attribute."}), 400
            
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        
        # Verify partition key exists in document
        pk_path = get_partition_key_path(container)
        pk_val = extract_partition_key_value(data, pk_path)
        if pk_val is None:
            clean_pk = pk_path.strip("/")
            return jsonify({
                "status": "error", 
                "message": f"Document must contain the partition key path attribute: '{clean_pk}'"
            }), 400
            
        # Execute Upsert with 429 retry
        res, _ = execute_with_429_retry(container.upsert_item, body=data)
        log_activity("cosmos_db", "UPSERT_DOCUMENT", f"{db_id}/{container_id}/{data.get('id')}", status="SUCCESS")
        return jsonify({"status": "success", "message": "Document saved successfully", "item": res})

    except Exception as e:
        traceback.print_exc()
        log_activity("cosmos_db", "UPSERT_DOCUMENT", f"{db_id}/{container_id}/{data.get('id', 'unknown')}", status="FAILED", details=str(e))
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/db/<db_id>/container/<container_id>/item/delete", methods=["POST"])
@write_permission_required
def api_delete_item(db_id, container_id):
    """Delete an item from Cosmos DB"""
    client = get_cosmos_client()
    
    try:
        data = request.get_json()
        item_id = data.get("id")
        partition_key = data.get("partition_key")
        
        if not item_id:
            return jsonify({"status": "error", "message": "Item ID is required."}), 400

        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        
        # Delete item with 429 retry
        execute_with_429_retry(container.delete_item, item=item_id, partition_key=partition_key)
        log_activity("cosmos_db", "DELETE_DOCUMENT", f"{db_id}/{container_id}/{item_id}", status="SUCCESS")
        return jsonify({"status": "success", "message": "Document deleted successfully."})

    except Exception as e:
        traceback.print_exc()
        log_activity("cosmos_db", "DELETE_DOCUMENT", f"{db_id}/{container_id}/{data.get('id', 'unknown')}", status="FAILED", details=str(e))
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/db/<db_id>/container/<container_id>/items/bulk-delete", methods=["POST"])
@write_permission_required
def api_bulk_delete_items(db_id, container_id):
    """Bulk delete a list of items from Cosmos DB with 429 rate limit backoff"""
    client = get_cosmos_client()
    try:
        data = request.get_json()
        if not data or "items" not in data:
            return jsonify({"status": "error", "message": "No items provided for deletion."}), 400
            
        items_to_delete = data.get("items", [])
        if not items_to_delete:
            return jsonify({"status": "error", "message": "Item list is empty."}), 400

        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        
        deleted_count = 0
        failed_count = 0
        total_retries = 0
        errors = []
        
        def delete_single_item(item_info):
            item_id = item_info.get("id")
            pk_val = item_info.get("partition_key")
            if not item_id:
                return False, 0, "Missing item ID"
            try:
                _, retries = execute_with_429_retry(container.delete_item, item=item_id, partition_key=pk_val)
                return True, retries, None
            except Exception as e:
                return False, 0, f"ID {item_id}: {str(e)}"
                
        max_workers = min(25, max(1, len(items_to_delete)))
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_item = {executor.submit(delete_single_item, item): item for item in items_to_delete}
            for future in as_completed(future_to_item):
                success, retries, err_msg = future.result()
                total_retries += retries
                if success:
                    deleted_count += 1
                else:
                    failed_count += 1
                    if len(errors) < 5 and err_msg:
                        errors.append(err_msg)
                        
        msg = f"Successfully deleted {deleted_count} document(s)."
        if total_retries > 0:
            msg += f" (Handled {total_retries} rate limit retries)"
        if failed_count > 0:
            msg += f" {failed_count} item(s) failed."
            
        log_activity("cosmos_db", "BULK_DELETE", f"{db_id}/{container_id}", status="SUCCESS" if deleted_count > 0 else "FAILED", details=f"Deleted: {deleted_count}, Failed: {failed_count}")
        return jsonify({
            "status": "success" if deleted_count > 0 else "error",
            "message": msg,
            "deleted_count": deleted_count,
            "failed_count": failed_count,
            "retries_429": total_retries,
            "errors": errors
        })
    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/db/<db_id>/container/<container_id>/import-async", methods=["POST"])
@write_permission_required
def api_async_import_items(db_id, container_id):
    """Initiates an asynchronous background bulk ingestion task with live progress tracking"""
    client = get_cosmos_client()
    
    file = request.files.get("file")
    if not file or file.filename == "":
        return jsonify({"status": "error", "message": "No file selected for import."}), 400
        
    filename = file.filename
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ["json", "jsonl", "ndjson", "csv", "xlsx"]:
        return jsonify({"status": "error", "message": "Unsupported format. Please upload JSON, JSONL, NDJSON, CSV, or XLSX."}), 400

    try:
        concurrency = request.form.get("concurrency", 100, type=int)
        concurrency = max(10, min(concurrency, 200))
        
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        pk_path = get_partition_key_path(container)
        
        task_id = str(uuid.uuid4())
        temp_file_path = os.path.join(TEMP_UPLOAD_DIR, f"{task_id}_{filename}")
        file.save(temp_file_path)
        
        # Estimate total records for quick progress estimation
        total_estimate = 0
        try:
            if ext in ["jsonl", "ndjson", "csv"]:
                with open(temp_file_path, "r", encoding="utf-8", errors="replace") as f:
                    for _ in f:
                        total_estimate += 1
                if ext == "csv" and total_estimate > 0:
                    total_estimate -= 1
        except Exception:
            total_estimate = 0

        with IMPORT_TASKS_LOCK:
            IMPORT_TASKS[task_id] = {
                "task_id": task_id,
                "db_id": db_id,
                "container_id": container_id,
                "filename": filename,
                "status": "starting",
                "total_estimate": total_estimate,
                "processed": 0,
                "successful": 0,
                "failed": 0,
                "retries_429": 0,
                "speed_per_sec": 0,
                "start_time": time.time(),
                "end_time": None,
                "errors": [],
                "cancel_requested": False
            }
            
        # Spawn daemon worker thread
        worker_thread = threading.Thread(
            target=run_bulk_import_worker,
            args=(task_id, temp_file_path, filename, db_id, container_id, pk_path, concurrency, client),
            daemon=True
        )
        worker_thread.start()
        
        log_activity("cosmos_db", "BULK_IMPORT_START", f"{db_id}/{container_id}", status="SUCCESS", details=f"File: {filename}, Concurrency: {concurrency}")
        return jsonify({
            "status": "success",
            "task_id": task_id,
            "message": "Bulk import job initiated successfully.",
            "total_estimate": total_estimate
        })

    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": f"Failed to start import task: {str(e)}"}), 500

@ui.route("/api/import-task/<task_id>", methods=["GET"])
@login_required
def api_get_import_task(task_id):
    """Fetch live progress metrics of a background bulk import task"""
    with IMPORT_TASKS_LOCK:
        task = IMPORT_TASKS.get(task_id)
        if not task:
            return jsonify({"status": "error", "message": "Task not found"}), 404
            
        elapsed = (task["end_time"] if task.get("end_time") else time.time()) - task["start_time"]
        
        return jsonify({
            "status": "success",
            "task": {
                "task_id": task["task_id"],
                "status": task["status"],
                "filename": task["filename"],
                "total_estimate": task["total_estimate"],
                "processed": task["processed"],
                "successful": task["successful"],
                "failed": task["failed"],
                "retries_429": task["retries_429"],
                "speed_per_sec": task["speed_per_sec"],
                "elapsed_seconds": round(elapsed, 1),
                "errors": task["errors"][:15],
                "error_message": task.get("error_message", None)
            }
        })

@ui.route("/api/import-task/<task_id>/cancel", methods=["POST"])
@write_permission_required
def api_cancel_import_task(task_id):
    """Requests graceful cancellation of a running bulk import task"""
    with IMPORT_TASKS_LOCK:
        task = IMPORT_TASKS.get(task_id)
        if not task:
            return jsonify({"status": "error", "message": "Task not found"}), 404
        task["cancel_requested"] = True
        task["status"] = "cancelled"
        task["end_time"] = time.time()
        log_activity("cosmos_db", "BULK_IMPORT_CANCEL", f"{task.get('db_id')}/{task.get('container_id')}", status="SUCCESS", details=f"Task: {task_id}")
        return jsonify({"status": "success", "message": "Cancellation requested."})

@ui.route("/api/db/<db_id>/container/<container_id>/empty", methods=["POST"])
@write_permission_required
def api_empty_container(db_id, container_id):
    """
    Empties all documents from a container by recreating it with identical schema/configuration
    or by bulk-deleting all documents.
    """
    client = get_cosmos_client()
    try:
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        
        # Read current container properties (partition key, indexing policy, default_ttl, etc.)
        properties = container.read()
        pk_info = properties.get("partitionKey", {})
        pk_paths = pk_info.get("paths", ["/id"])
        pk_path = pk_paths[0] if pk_paths else "/id"
        indexing_policy = properties.get("indexingPolicy")
        default_ttl = properties.get("defaultTtl")
        unique_key_policy = properties.get("uniqueKeyPolicy")
        
        # Method 1: Drop & Recreate container (Instant deletion of millions of documents with 0 RU cost per item)
        try:
            db_client.delete_container(container_id)
            
            create_kwargs = {
                "id": container_id,
                "partition_key": PartitionKey(path=pk_path)
            }
            if indexing_policy:
                create_kwargs["indexing_policy"] = indexing_policy
            if default_ttl is not None:
                create_kwargs["default_ttl"] = default_ttl
            if unique_key_policy:
                create_kwargs["unique_key_policy"] = unique_key_policy
                
            db_client.create_container(**create_kwargs)
            log_activity("cosmos_db", "EMPTY_CONTAINER", f"{db_id}/{container_id}", status="SUCCESS", details="Container recreated with original schema")
            return jsonify({
                "status": "success",
                "message": f"Container '{container_id}' was emptied successfully (recreated with original schema)."
            })
        except Exception as recreate_err:
            print(f"Container recreate notice, falling back to document truncate: {recreate_err}")
            # Method 2: Fallback query and bulk delete
            items = list(container.query_items("SELECT c.id FROM c", enable_cross_partition_query=True))
            deleted_count = 0
            for it in items:
                pk_val = extract_partition_key_value(it, pk_path)
                try:
                    execute_with_429_retry(container.delete_item, item=it["id"], partition_key=pk_val)
                    deleted_count += 1
                except Exception:
                    pass
            log_activity("cosmos_db", "EMPTY_CONTAINER", f"{db_id}/{container_id}", status="SUCCESS", details=f"Truncated {deleted_count} items")
            return jsonify({
                "status": "success",
                "message": f"Emptied {deleted_count} document(s) from container '{container_id}'."
            })
            
    except Exception as e:
        traceback.print_exc()
        log_activity("cosmos_db", "EMPTY_CONTAINER", f"{db_id}/{container_id}", status="FAILED", details=str(e))
        return jsonify({"status": "error", "message": f"Failed to empty container: {str(e)}"}), 500

@ui.route("/api/db/<db_id>/container/<container_id>/export-async", methods=["POST"])
@login_required
def api_async_export_items(db_id, container_id):
    """Initiates an asynchronous background export task with live progress tracking"""
    client = get_cosmos_client()

    try:
        data = request.get_json() or {}
        format_type = data.get("format", "jsonl").lower()
        if format_type not in ["json", "jsonl", "ndjson", "csv", "xlsx"]:
            format_type = "jsonl"

        search_mode = data.get("search_mode", "simple")
        search_query = data.get("search_query", "")

        q_meta = parse_and_build_query(search_mode, search_query, offset=0, limit=10000000)
        export_sql = q_meta["items_sql"]
        # Strip pagination offset/limit to export full dataset
        if "OFFSET" in export_sql.upper():
            offset_pos = export_sql.upper().rfind("OFFSET")
            export_sql = export_sql[:offset_pos].strip()

        task_id = str(uuid.uuid4())
        ext = "xlsx" if format_type == "xlsx" else ("csv" if format_type == "csv" else ("json" if format_type == "json" else "jsonl"))
        filename = f"cosmos_{container_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"
        temp_file_path = os.path.join(TEMP_EXPORT_DIR, f"{task_id}_{filename}")

        # Total count estimate
        total_estimate = data.get("total_estimate", 0)

        with EXPORT_TASKS_LOCK:
            EXPORT_TASKS[task_id] = {
                "task_id": task_id,
                "db_id": db_id,
                "container_id": container_id,
                "filename": filename,
                "file_path": temp_file_path,
                "format": format_type,
                "status": "starting",
                "total_estimate": total_estimate,
                "processed": 0,
                "retries_429": 0,
                "speed_per_sec": 0,
                "start_time": time.time(),
                "end_time": None,
                "cancel_requested": False
            }

        worker_thread = threading.Thread(
            target=run_export_worker,
            args=(task_id, temp_file_path, filename, format_type, db_id, container_id, export_sql, q_meta["params"], client),
            daemon=True
        )
        worker_thread.start()
        log_activity("cosmos_db", "EXPORT_DATA", f"{db_id}/{container_id}", status="SUCCESS", details=f"Format: {format_type}, File: {filename}")

        return jsonify({
            "status": "success",
            "task_id": task_id,
            "message": "Export job initiated successfully.",
            "filename": filename
        })

    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": f"Failed to start export: {str(e)}"}), 500

@ui.route("/api/export-task/<task_id>", methods=["GET"])
@login_required
def api_get_export_task(task_id):
    """Fetch live progress metrics of a background export task"""
    with EXPORT_TASKS_LOCK:
        task = EXPORT_TASKS.get(task_id)
        if not task:
            return jsonify({"status": "error", "message": "Task not found"}), 404

        elapsed = (task["end_time"] if task.get("end_time") else time.time()) - task["start_time"]

        return jsonify({
            "status": "success",
            "task": {
                "task_id": task["task_id"],
                "status": task["status"],
                "filename": task["filename"],
                "total_estimate": task["total_estimate"],
                "processed": task["processed"],
                "retries_429": task["retries_429"],
                "speed_per_sec": task["speed_per_sec"],
                "elapsed_seconds": round(elapsed, 1),
                "error_message": task.get("error_message", None)
            }
        })

@ui.route("/api/export-task/<task_id>/cancel", methods=["POST"])
@login_required
def api_cancel_export_task(task_id):
    """Requests cancellation of a running export task"""
    with EXPORT_TASKS_LOCK:
        task = EXPORT_TASKS.get(task_id)
        if not task:
            return jsonify({"status": "error", "message": "Task not found"}), 404
        task["cancel_requested"] = True
        task["status"] = "cancelled"
        task["end_time"] = time.time()
        log_activity("cosmos_db", "EXPORT_CANCEL", f"{task.get('db_id')}/{task.get('container_id')}", status="SUCCESS", details=f"Task: {task_id}")
        return jsonify({"status": "success", "message": "Export cancellation requested."})

@ui.route("/api/export-task/<task_id>/download", methods=["GET"])
@login_required
def api_download_export_task(task_id):
    """Download the completed exported file"""
    with EXPORT_TASKS_LOCK:
        task = EXPORT_TASKS.get(task_id)
        if not task or task.get("status") != "completed":
            flash("Export file not ready or task not found.", "warning")
            return redirect(url_for("ui.dashboard"))

        file_path = task["file_path"]
        filename = task["filename"]

    if not os.path.exists(file_path):
        flash("Export file does not exist on server.", "danger")
        return redirect(url_for("ui.dashboard"))

    mimetype = "application/octet-stream"
    if filename.endswith(".json"):
        mimetype = "application/json"
    elif filename.endswith(".jsonl") or filename.endswith(".ndjson"):
        mimetype = "application/x-ndjson"
    elif filename.endswith(".csv"):
        mimetype = "text/csv"
    elif filename.endswith(".xlsx"):
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return send_file(
        file_path,
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename
    )

@ui.route("/db/<db_id>/container/<container_id>/import", methods=["POST"])
@write_permission_required
def import_items(db_id, container_id):
    """Synchronous import fallback handling JSON, JSONL, CSV, and XLSX with 429 retries"""
    client = get_cosmos_client()
    
    file = request.files.get("file")
    if not file or file.filename == "":
        flash("No file selected for import.", "warning")
        return redirect(url_for("ui.container_view", db_id=db_id, container_id=container_id))

    try:
        filename = file.filename
        temp_path = os.path.join(TEMP_UPLOAD_DIR, f"sync_{uuid.uuid4()}_{filename}")
        file.save(temp_path)
        
        db_client = client.get_database_client(db_id)
        container = db_client.get_container_client(container_id)
        pk_path = get_partition_key_path(container)
        clean_pk = pk_path.strip("/") if pk_path else "id"

        imported_count = 0
        skipped_count = 0
        total_429_retries = 0
        errors = []

        try:
            for idx, doc in stream_file_records(temp_path, filename):
                if "_parse_error" in doc:
                    skipped_count += 1
                    errors.append(f"Row {idx}: {doc['_parse_error']}")
                    continue
                try:
                    if "id" not in doc or not str(doc["id"]).strip():
                        doc["id"] = str(uuid.uuid4())
                    pk_val = extract_partition_key_value(doc, pk_path)
                    if pk_val is None:
                        doc[clean_pk] = "imported"
                    _, retries = execute_with_429_retry(container.upsert_item, body=doc, max_retries=10)
                    total_429_retries += retries
                    imported_count += 1
                except Exception as item_err:
                    skipped_count += 1
                    if len(errors) < 5:
                        errors.append(f"Row {idx}: {str(item_err)}")
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

        msg = f"Import Summary: {imported_count} documents imported successfully."
        if total_429_retries > 0:
            msg += f" (Handled {total_429_retries} rate limit 429 retries)"
        if skipped_count > 0:
            msg += f" {skipped_count} items failed. Sample errors: {', '.join(errors[:3])}"
            flash(msg, "warning")
        else:
            flash(msg, "success")
        log_activity("cosmos_db", "SYNC_IMPORT", f"{db_id}/{container_id}", status="SUCCESS" if imported_count > 0 else "FAILED", details=f"Imported: {imported_count}, Skipped: {skipped_count}")

    except Exception as e:
        traceback.print_exc()
        flash(f"Import process failed: {str(e)}", "danger")

    return redirect(url_for("ui.container_view", db_id=db_id, container_id=container_id))

@ui.route("/api/provision", methods=["POST"])
@write_permission_required
def api_provision():
    """Provision a new Database and optionally a Container inside it"""
    client = get_cosmos_client()
    try:
        data = request.get_json()
        if not data:
            return jsonify({"status": "error", "message": "No data provided"}), 400
            
        db_id = data.get("db_id", "").strip()
        container_id = data.get("container_id", "").strip()
        partition_key = data.get("partition_key", "").strip()
        
        if not db_id:
            return jsonify({"status": "error", "message": "Database ID is required."}), 400
            
        existing_dbs = {db["id"] for db in client.list_databases()}
        
        if container_id:
            if db_id in existing_dbs:
                db_client = client.get_database_client(db_id)
                existing_containers = {c["id"] for c in db_client.list_containers()}
                if container_id in existing_containers:
                    return jsonify({"status": "error", "message": f"Container '{container_id}' already exists in database '{db_id}'."}), 400
            
            client.create_database_if_not_exists(id=db_id)
            
            if not partition_key:
                partition_key = "/id"
            if not partition_key.startswith("/"):
                partition_key = "/" + partition_key
                
            db_client = client.get_database_client(db_id)
            db_client.create_container_if_not_exists(
                id=container_id,
                partition_key=PartitionKey(path=partition_key)
            )
            msg = f"Successfully provisioned database '{db_id}' and container '{container_id}'."
        else:
            if db_id in existing_dbs:
                return jsonify({"status": "error", "message": f"Database '{db_id}' already exists."}), 400
                
            client.create_database_if_not_exists(id=db_id)
            msg = f"Successfully created database '{db_id}'."
            
        log_activity("cosmos_db", "PROVISION_STRUCTURE", f"{db_id}/{container_id}" if container_id else db_id, status="SUCCESS")
        return jsonify({
            "status": "success", 
            "message": msg
        })
    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/db/<db_id>/delete", methods=["POST"])
@admin_required
def api_delete_database(db_id):
    """Delete a Database from Cosmos DB (Admin Only)"""
    client = get_cosmos_client()
    try:
        client.delete_database(db_id)
        log_activity("cosmos_db", "DELETE_DATABASE", db_id, status="SUCCESS")
        return jsonify({"status": "success", "message": f"Database '{db_id}' deleted successfully."})
    except Exception as e:
        traceback.print_exc()
        log_activity("cosmos_db", "DELETE_DATABASE", db_id, status="FAILED", details=str(e))
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/db/<db_id>/container/<container_id>/delete", methods=["POST"])
@admin_required
def api_delete_container(db_id, container_id):
    """Delete a Container from a Database (Admin Only)"""
    client = get_cosmos_client()
    try:
        db_client = client.get_database_client(db_id)
        db_client.delete_container(container_id)
        log_activity("cosmos_db", "DELETE_CONTAINER", f"{db_id}/{container_id}", status="SUCCESS")
        return jsonify({"status": "success", "message": f"Container '{container_id}' deleted successfully from database '{db_id}'."})
    except Exception as e:
        traceback.print_exc()
        log_activity("cosmos_db", "DELETE_CONTAINER", f"{db_id}/{container_id}", status="FAILED", details=str(e))
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/api/bulk-check", methods=["POST"])
@write_permission_required
def api_bulk_check():
    """Parse bulk upload file and check for existing databases"""
    client = get_cosmos_client()
    
    file = request.files.get("file")
    if not file or file.filename == "":
        return jsonify({"status": "error", "message": "No file provided"}), 400
        
    try:
        filename = file.filename.lower()
        rows = []
        headers = []
        
        file_bytes = file.stream.read()
        file_io = io.BytesIO(file_bytes)

        if filename.endswith(".csv"):
            stream = io.StringIO(file_bytes.decode("utf-8"), newline=None)
            reader = csv.reader(stream)
            headers = next(reader, None)
            if headers:
                for r in reader:
                    rows.append(r)
        elif filename.endswith(".xlsx"):
            wb = load_workbook(file_io, read_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = next(rows_iter, None)
            if headers:
                for r in rows_iter:
                    rows.append(r)
        else:
            return jsonify({"status": "error", "message": "Unsupported file format"}), 400
            
        if not headers:
            return jsonify({"status": "error", "message": "File is empty or missing headers"}), 400
            
        db_idx = -1
        for h_idx, h in enumerate(headers):
            if h is None:
                continue
            h_str = str(h).strip().lower()
            if h_str in ["db_name", "database_name", "database", "db", "dbname"]:
                db_idx = h_idx
        
        if db_idx == -1 and len(headers) > 0:
            db_idx = 0
            
        uploaded_dbs = set()
        for row in rows:
            if not row or all(v is None for v in row):
                continue
            db_name = str(row[db_idx]).strip() if db_idx < len(row) and row[db_idx] is not None else ""
            if db_name:
                uploaded_dbs.add(db_name)
                
        existing_cosmos_dbs = {db["id"] for db in client.list_databases() if db["id"] != SYSTEM_DB}
        overlap = list(uploaded_dbs.intersection(existing_cosmos_dbs))
        
        return jsonify({"status": "success", "existing_dbs": overlap})
        
    except Exception as e:
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500

@ui.route("/bulk-create", methods=["POST"])
@write_permission_required
def bulk_create_dbs_containers():
    """Bulk provision databases and containers from CSV or Excel (.xlsx) file"""
    client = get_cosmos_client()
    
    file = request.files.get("file")
    if not file or file.filename == "":
        flash("No file selected for bulk creation.", "warning")
        return redirect(url_for("ui.dashboard"))
        
    try:
        filename = file.filename.lower()
        rows = []
        headers = []
        
        file_bytes = file.stream.read()
        file_io = io.BytesIO(file_bytes)

        if filename.endswith(".csv"):
            stream = io.StringIO(file_bytes.decode("utf-8"), newline=None)
            reader = csv.reader(stream)
            headers = next(reader, None)
            if headers:
                for r in reader:
                    rows.append(r)
        elif filename.endswith(".xlsx"):
            wb = load_workbook(file_io, read_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = next(rows_iter, None)
            if headers:
                for r in rows_iter:
                    rows.append(r)
        else:
            flash("Unsupported file format. Please upload a CSV or XLSX file.", "danger")
            return redirect(url_for("ui.dashboard"))
            
        if not headers:
            flash("Uploaded file is empty or missing headers.", "danger")
            return redirect(url_for("ui.dashboard"))
            
        # Standardize headers to find indices
        db_idx, container_idx, pk_idx = -1, -1, -1
        for h_idx, h in enumerate(headers):
            if h is None:
                continue
            h_str = str(h).strip().lower()
            if h_str in ["db_name", "database_name", "database", "db", "dbname"]:
                db_idx = h_idx
            elif h_str in ["container_name", "container", "collection", "containername"]:
                container_idx = h_idx
            elif h_str in ["partition_key", "partition_path", "partitionkey", "pk", "partitionkeypath"]:
                pk_idx = h_idx
                
        # Robust positions fallback if header names aren't matched
        if db_idx == -1 and len(headers) > 0:
            db_idx = 0
        if container_idx == -1 and len(headers) > 1:
            container_idx = 1
        if pk_idx == -1 and len(headers) > 2:
            pk_idx = 2
            
        created_dbs = set()
        created_containers = []
        errors = []
        
        for idx, row in enumerate(rows):
            if not row or all(v is None for v in row):
                continue
                
            try:
                db_name = str(row[db_idx]).strip() if db_idx < len(row) and row[db_idx] is not None else ""
                container_name = str(row[container_idx]).strip() if container_idx < len(row) and row[container_idx] is not None else ""
                partition_key = str(row[pk_idx]).strip() if (pk_idx != -1 and pk_idx < len(row) and row[pk_idx] is not None) else "/id"
                
                if not db_name or not container_name:
                    continue
                    
                mode = request.form.get("mode", "merge")
                
                # Create Database if not exists
                if mode == "overwrite" and db_name not in created_dbs:
                    try:
                        client.get_database_client(db_name).read()
                        client.delete_database(db_name)
                    except exceptions.CosmosResourceNotFoundError:
                        pass

                client.create_database_if_not_exists(id=db_name)
                created_dbs.add(db_name)
                
                # Clean Partition Key
                pk_path = partition_key if partition_key else "/id"
                if not pk_path.startswith("/"):
                    pk_path = "/" + pk_path
                    
                # Create Container if not exists
                db_client = client.get_database_client(db_name)
                db_client.create_container_if_not_exists(
                    id=container_name,
                    partition_key=PartitionKey(path=pk_path)
                )
                created_containers.append(f"{db_name}/{container_name}")
                
            except Exception as row_err:
                errors.append(f"Row {idx+2}: {str(row_err)}")
                
        msg = f"Bulk Process Complete! Verified/Created {len(created_dbs)} databases and {len(created_containers)} containers."
        if errors:
            msg += f" Encountered {len(errors)} errors. Sample errors: {', '.join(errors[:3])}"
            flash(msg, "warning")
        else:
            flash(msg, "success")
        log_activity("cosmos_db", "BULK_PROVISION", file.filename, status="SUCCESS" if len(created_containers) > 0 else "FAILED", details=f"Databases: {len(created_dbs)}, Containers: {len(created_containers)}")
            
    except Exception as e:
        traceback.print_exc()
        flash(f"Bulk creation failed: {str(e)}", "danger")
        
    return redirect(url_for("ui.dashboard"))

# Register blueprint
app.register_blueprint(ui)

# ---------- Run App ----------
if __name__ == "__main__":
    # Local dev server runs on 5001 to prevent conflicts
    app.run(host="0.0.0.0", port=5001, debug=True)